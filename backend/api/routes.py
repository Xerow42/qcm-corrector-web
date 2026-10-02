from __future__ import annotations
from pathlib import Path
from datetime import datetime
from email.mime.image import MIMEImage
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import csv
import io
import json
import os
import re
import smtplib
import ssl
import subprocess
import tempfile
import unicodedata
from urllib.parse import quote
from fastapi import APIRouter, File, Query, UploadFile, HTTPException, Body, Response, Form, Header
from fastapi.responses import FileResponse
import cv2
from db import connect, load_local_env, ping_database
from preprocessing.opencv_pipeline import preprocess_sheet
from segmentation.grid_extractor import extract_rois, compute_boxes
from classification.pixel_threshold import black_ratio
from classification.cnn_classifier import classify_cnn
from scoring import score_answers
from pdf_generator import generate_questionnaire_pdf, generate_statements_pdf

router = APIRouter()
BASE_DIR = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = BASE_DIR / "configs" / "template_60q.json"
MODEL_PATH = BASE_DIR / "models" / "model.h5"
ANSWER_KEY_PATH = BASE_DIR / "data" / "answer_key.json"
TEMPLATE_DIR = BASE_DIR / "data" / "templates"
TEMPLATE_PDF_PATH = TEMPLATE_DIR / "template.pdf"
DEBUG_DIR = BASE_DIR / "data" / "samples"
SUBMISSION_DIR = BASE_DIR / "data" / "submissions"
DEBUG_UPLOAD_BASENAME = "last_upload"
DEBUG_OVERLAY_PATH = DEBUG_DIR / "debug_overlay_last.jpg"
load_local_env()
IA_API_BASE_URL = os.environ.get("IA_API_BASE_URL", "http://127.0.0.1:8000").rstrip("/")
IA_ANALYZE_PATH = os.environ.get("IA_ANALYZE_PATH", "/analyze")
IA_TIMEOUT_SECONDS = float(os.environ.get("IA_TIMEOUT_SECONDS", "60"))
AI_RESULTS_API_KEY = os.environ.get("AI_RESULTS_API_KEY", "").strip()
WEB_BASE_URL = os.environ.get("WEB_BASE_URL", "http://localhost:3000").rstrip("/")
SMTP_HOST = os.environ.get("SMTP_HOST", "").strip()
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587") or "587")
SMTP_USER = os.environ.get("SMTP_USER", "").strip()
SMTP_PASSWORD = os.environ.get("SMTP_PASSWORD", "")
SMTP_FROM = os.environ.get("SMTP_FROM", SMTP_USER or "no-reply@junia.ma").strip()
SMTP_USE_TLS = os.environ.get("SMTP_USE_TLS", "true").strip().lower() not in {"0", "false", "no"}
LOGO_PATH = BASE_DIR.parent / "web" / "public" / "logo-junia.png"


def _find_browser_executable() -> str | None:
    candidates = [
        os.environ.get("CHROME_PATH"),
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    return None


def _render_web_page_pdf(path: str, filename_hint: str) -> bytes:
    browser = _find_browser_executable()
    if not browser:
        raise HTTPException(status_code=500, detail="Chrome ou Edge introuvable pour generer le PDF HTML.")

    url = f"{WEB_BASE_URL}{path}"
    with tempfile.TemporaryDirectory(prefix="qcm_pdf_") as temp_dir:
        pdf_path = Path(temp_dir) / filename_hint
        profile_dir = Path(temp_dir) / "profile"
        cmd = [
            browser,
            "--headless=new",
            "--disable-gpu",
            "--disable-dev-shm-usage",
            "--no-sandbox",
            "--no-pdf-header-footer",
            "--print-to-pdf-no-header",
            "--virtual-time-budget=5000",
            f"--user-data-dir={profile_dir}",
            f"--print-to-pdf={pdf_path}",
            url,
        ]
        try:
            completed = subprocess.run(cmd, check=False, capture_output=True, text=True, timeout=45)
        except subprocess.TimeoutExpired as exc:
            raise HTTPException(status_code=504, detail="Generation PDF trop longue. Verifiez que le front Next.js est lance.") from exc

        if completed.returncode != 0 or not pdf_path.exists() or pdf_path.stat().st_size == 0:
            detail = (completed.stderr or completed.stdout or "").strip()
            raise HTTPException(
                status_code=500,
                detail=f"Generation PDF HTML impossible. Verifiez que {WEB_BASE_URL} est accessible. {detail}",
            )
        return pdf_path.read_bytes()


def _smtp_settings() -> dict:
    load_local_env(force=True)
    host = os.environ.get("SMTP_HOST", "").strip()
    user = os.environ.get("SMTP_USER", "").strip()
    return {
        "host": host,
        "port": int(os.environ.get("SMTP_PORT", "587") or "587"),
        "user": user,
        "password": os.environ.get("SMTP_PASSWORD", ""),
        "from": os.environ.get("SMTP_FROM", user or "no-reply@junia.ma").strip(),
        "use_tls": os.environ.get("SMTP_USE_TLS", "true").strip().lower() not in {"0", "false", "no"},
    }


def _normalize_student_number(raw: object) -> str:
    """Keep the six-digit student identifier stable, including leading zeroes."""
    text = str(raw or "").strip()
    if not text:
        return ""
    digits = re.sub(r"\D", "", text)
    if digits:
        return digits.zfill(6) if len(digits) <= 6 else digits
    return text


def load_template() -> dict:
    with open(TEMPLATE_PATH, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def get_template_image_path() -> Path | None:
    for ext in (".jpg", ".jpeg", ".png"):
        p = TEMPLATE_DIR / f"template{ext}"
        if p.exists():
            return p
    return None


def get_template_pdf_path() -> Path | None:
    if TEMPLATE_PDF_PATH.exists():
        return TEMPLATE_PDF_PATH
    return None

def get_template_path() -> tuple[str | None, Path | None]:
    pdf = get_template_pdf_path()
    if pdf:
        return "pdf", pdf
    image = get_template_image_path()
    if image:
        return "image", image
    return None, None


def _save_last_upload(payload: bytes, filename: str | None) -> Path | None:
    ext = (Path(filename or "").suffix or ".jpg").lower()
    if ext not in (".jpg", ".jpeg", ".png"):
        ext = ".jpg"
    DEBUG_DIR.mkdir(parents=True, exist_ok=True)
    for old_ext in (".jpg", ".jpeg", ".png"):
        old = DEBUG_DIR / f"{DEBUG_UPLOAD_BASENAME}{old_ext}"
        if old.exists():
            old.unlink()
    path = DEBUG_DIR / f"{DEBUG_UPLOAD_BASENAME}{ext}"
    with open(path, "wb") as f:
        f.write(payload)
    return path


def _get_last_upload_path() -> Path | None:
    for ext in (".jpg", ".jpeg", ".png"):
        path = DEBUG_DIR / f"{DEBUG_UPLOAD_BASENAME}{ext}"
        if path.exists():
            return path
    return None


def _extract_numeric_id(raw_id: str | int) -> int:
    text = str(raw_id).strip()
    if text.isdigit():
        return int(text)
    match = re.search(r"(\d+)$", text)
    if match:
        return int(match.group(1))
    raise HTTPException(status_code=400, detail=f"Identifiant invalide: {raw_id}")


def _resolve_viewer(
    role: str | None = None,
    user_id: str | int | None = None,
    email: str | None = None,
) -> dict | None:
    """Resolve the optional web viewer used to scope professor dashboards."""
    raw_role = str(role or "").strip().lower()
    raw_email = str(email or "").strip().lower()
    raw_user_id = str(user_id or "").strip()
    if not raw_role and not raw_user_id and not raw_email:
        return None
    if raw_role not in {"admin", "enseignant"}:
        raise HTTPException(status_code=400, detail="Role utilisateur invalide.")

    clauses: list[str] = []
    params: list[object] = []
    if raw_user_id:
        clauses.append("id_utilisateur = %s")
        params.append(_extract_numeric_id(raw_user_id))
    if raw_email:
        clauses.append("lower(email) = %s")
        params.append(raw_email)
    if not clauses:
        raise HTTPException(status_code=400, detail="Utilisateur connecte introuvable: id ou email requis.")

    query = f"""
        SELECT id_utilisateur, email, nom, prenom, role::text AS role, actif
        FROM utilisateur
        WHERE ({' OR '.join(clauses)})
          AND role = %s
          AND actif = TRUE
        LIMIT 1
    """
    params.append(raw_role)
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(query, params)
            row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=403, detail="Utilisateur non autorise ou inactif.")
    return {
        "id": int(row["id_utilisateur"]),
        "email": row["email"],
        "first_name": row["prenom"],
        "last_name": row["nom"],
        "role": row["role"],
    }


def _is_teacher_viewer(viewer: dict | None) -> bool:
    return bool(viewer and viewer.get("role") == "enseignant")


def _assert_teacher_can_access_session(viewer: dict | None, session_ctx: dict) -> None:
    if not _is_teacher_viewer(viewer):
        return
    teacher_id = int(viewer["id"])
    allowed_ids = {
        int(value)
        for value in (
            session_ctx.get("id_enseignant"),
            session_ctx.get("id_createur"),
            session_ctx.get("id_matiere_enseignant"),
        )
        if value is not None
    }
    if teacher_id not in allowed_ids:
        raise HTTPException(status_code=403, detail="Session reservee a l'enseignant rattache a la matiere.")


def _assert_teacher_can_access_qcm(viewer: dict | None, qcm_detail: dict) -> None:
    if not _is_teacher_viewer(viewer):
        return
    teacher_id = int(viewer["id"])
    qcm_teacher_id = qcm_detail.get("teacher", {}).get("id")
    if qcm_teacher_id is None or int(qcm_teacher_id) != teacher_id:
        raise HTTPException(status_code=403, detail="QCM reserve a l'enseignant rattache a la matiere.")


def _fetch_session_context(session_identifier: str) -> dict:
    session_id = _extract_numeric_id(session_identifier)
    query = """
        SELECT
            s.id_session,
            s.libelle,
            s.statut::text AS statut,
            s.email_status::text AS email_status,
            s.id_qcm,
            q.code AS qcm_code,
            q.intitule AS qcm_title,
            q.type_qcm::text AS qcm_type,
            q.statut::text AS qcm_status,
            q.date_examen,
            q.id_createur,
            m.id_matiere,
            m.id_enseignant AS id_matiere_enseignant,
            s.id_enseignant,
            u.nom AS enseignant_nom,
            u.prenom AS enseignant_prenom,
            s.id_classe,
            c.nom AS classe_nom,
            c.niveau AS classe_niveau,
            s.id_concours,
            co.intitule AS concours_intitule
        FROM session_correction s
        JOIN qcm q ON q.id_qcm = s.id_qcm
        JOIN matiere m ON m.id_matiere = q.id_matiere
        JOIN utilisateur u ON u.id_utilisateur = s.id_enseignant
        LEFT JOIN classe c ON c.id_classe = s.id_classe
        LEFT JOIN concours co ON co.id_concours = s.id_concours
        WHERE s.id_session = %s
    """
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(query, (session_id,))
            row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail=f"Session introuvable: {session_identifier}")
    return row


def _fetch_qcm_questions(qcm_id: int) -> list[dict]:
    query = """
        SELECT
            q.id_question,
            q.numero,
            q.intitule AS question_text,
            q.ponderation,
            q.type_question::text AS type_question,
            r.id_reponse,
            r.label,
            r.intitule AS reponse_text,
            r.est_correcte
        FROM question q
        LEFT JOIN reponse r ON r.id_question = q.id_question
        WHERE q.id_qcm = %s
        ORDER BY q.numero, r.label
    """
    grouped: dict[int, dict] = {}
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(query, (qcm_id,))
            rows = cur.fetchall()
    for row in rows:
        question_id = int(row["id_question"])
        item = grouped.setdefault(
            question_id,
            {
                "id_question": question_id,
                "numero": int(row["numero"]),
                "intitule": row["question_text"],
                "ponderation": float(row["ponderation"]),
                "type_question": row["type_question"],
                "choices": {},
                "expected": [],
            },
        )
        if row.get("id_reponse") is not None:
            label = str(row["label"]).upper()
            item["choices"][label] = {
                "id_reponse": int(row["id_reponse"]),
                "intitule": row["reponse_text"],
                "est_correcte": bool(row["est_correcte"]),
            }
            if row["est_correcte"]:
                item["expected"].append(label)
    return sorted(grouped.values(), key=lambda entry: entry["numero"])


def _resolve_qcm_id(cur, raw_id: str | int) -> int:
    text = str(raw_id).strip()
    if text.isdigit():
        return int(text)
    if text.lower().startswith("qcm-") and text[4:].isdigit():
        return int(text[4:])
    cur.execute("SELECT id_qcm FROM qcm WHERE code = %s", (text,))
    row = cur.fetchone()
    if row:
        return int(row["id_qcm"])
    raise HTTPException(status_code=404, detail=f"QCM introuvable: {raw_id}")


def _fetch_qcm_detail(qcm_identifier: str | int) -> dict:
    query = """
        SELECT
            q.id_qcm,
            q.code,
            q.intitule,
            q.type_qcm::text AS type_qcm,
            q.statut::text AS statut,
            q.date_examen,
            q.id_createur,
            u.prenom AS teacher_first_name,
            u.nom AS teacher_last_name,
            u.email AS teacher_email,
            m.id_matiere,
            m.intitule AS subject_name,
            m.coefficient,
            c.id_classe,
            c.nom AS class_name,
            c.niveau AS class_level,
            c.annee_scolaire,
            q_totals.max_score
        FROM qcm q
        JOIN utilisateur u ON u.id_utilisateur = q.id_createur
        JOIN matiere m ON m.id_matiere = q.id_matiere
        LEFT JOIN classe c ON c.id_classe = q.id_classe
        LEFT JOIN (
            SELECT id_qcm, COALESCE(SUM(ponderation), 0) AS max_score
            FROM question
            GROUP BY id_qcm
        ) q_totals ON q_totals.id_qcm = q.id_qcm
        WHERE q.id_qcm = %s
    """
    with connect() as conn:
        with conn.cursor() as cur:
            qcm_id = _resolve_qcm_id(cur, qcm_identifier)
            cur.execute(query, (qcm_id,))
            row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail=f"QCM introuvable: {qcm_identifier}")

    questions = _fetch_qcm_questions(int(row["id_qcm"]))
    return {
        "id": f"qcm-{int(row['id_qcm'])}",
        "code": row["code"],
        "title": row["intitule"],
        "type": row["type_qcm"],
        "status": row["statut"],
        "exam_date": row.get("date_examen").isoformat() if row.get("date_examen") else "",
        "max_score": float(row.get("max_score") or 0),
        "teacher": {
            "id": int(row["id_createur"]),
            "first_name": row.get("teacher_first_name"),
            "last_name": row.get("teacher_last_name"),
            "email": row.get("teacher_email"),
        },
        "subject": {
            "id": int(row["id_matiere"]),
            "name": row.get("subject_name"),
            "coefficient": float(row.get("coefficient") or 1),
        },
        "class": {
            "id": int(row["id_classe"]) if row.get("id_classe") else None,
            "name": row.get("class_name"),
            "level": row.get("class_level"),
            "academic_year": row.get("annee_scolaire"),
        },
        "questions": [
            {
                "id": int(question["id_question"]),
                "order": int(question["numero"]),
                "prompt": question["intitule"],
                "points": float(question["ponderation"]),
                "type": "multiple" if question["type_question"] == "multiple" else "single",
                "options": [
                    {
                        "key": label,
                        "text": choice["intitule"],
                        "isCorrect": bool(choice["est_correcte"]),
                    }
                    for label, choice in sorted(question["choices"].items())
                ],
            }
            for question in questions
        ],
    }


def _qcm_status_from_web(status: object) -> str:
    mapping = {
        "draft": "brouillon",
        "published": "publie",
        "archived": "archive",
        "brouillon": "brouillon",
        "publie": "publie",
        "archive": "archive",
    }
    return mapping.get(str(status or "draft").strip().lower(), "brouillon")


def _qcm_type_from_web(audience: object) -> str:
    return "concours" if str(audience or "").strip().lower() == "contest" else "classe"


def _require_text(value: object, field_name: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise HTTPException(status_code=400, detail=f"Champ obligatoire manquant: {field_name}")
    return text


def _normalize_web_question(raw_question: object, index: int) -> dict:
    if not isinstance(raw_question, dict):
        raise HTTPException(status_code=400, detail=f"Question {index + 1} invalide.")

    prompt = _require_text(raw_question.get("prompt") or raw_question.get("intitule"), f"questions[{index}].prompt")
    points = float(raw_question.get("points") or raw_question.get("ponderation") or 1)
    if points <= 0:
        raise HTTPException(status_code=400, detail=f"La ponderation de la question {index + 1} doit etre positive.")

    raw_options = raw_question.get("options")
    if not isinstance(raw_options, list) or not raw_options:
        raise HTTPException(status_code=400, detail=f"La question {index + 1} doit contenir des propositions.")

    options: list[dict] = []
    correct_count = 0
    for option_index, raw_option in enumerate(raw_options):
        if not isinstance(raw_option, dict):
            raise HTTPException(status_code=400, detail=f"Proposition invalide pour la question {index + 1}.")
        label = str(raw_option.get("key") or raw_option.get("label") or "").strip().upper()
        if label not in {"A", "B", "C", "D"}:
            raise HTTPException(status_code=400, detail=f"Label invalide pour la question {index + 1}: {label or option_index + 1}")
        is_correct = bool(raw_option.get("isCorrect") or raw_option.get("est_correcte"))
        correct_count += 1 if is_correct else 0
        options.append(
            {
                "label": label,
                "text": str(raw_option.get("text") or raw_option.get("intitule") or label).strip() or label,
                "is_correct": is_correct,
            }
        )

    if correct_count == 0:
        raise HTTPException(status_code=400, detail=f"La question {index + 1} doit avoir au moins une bonne reponse.")

    question_type = "multiple" if correct_count > 1 or raw_question.get("type") == "multiple" else "unique"
    return {
        "numero": int(raw_question.get("order") or index + 1),
        "prompt": prompt,
        "points": points,
        "type_question": question_type,
        "options": sorted(options, key=lambda option: option["label"]),
    }


def _resolve_teacher_id(cur, payload: dict) -> int:
    teacher_email = str(payload.get("teacher_email") or "").strip().lower()
    if teacher_email:
        cur.execute(
            """
            SELECT id_utilisateur
            FROM utilisateur
            WHERE lower(email) = %s AND role = 'enseignant'
            """,
            (teacher_email,),
        )
        row = cur.fetchone()
        if row:
            return int(row["id_utilisateur"])

    raw_teacher_id = payload.get("teacher_id") or payload.get("teacherId")
    if raw_teacher_id:
        try:
            teacher_id = _extract_numeric_id(raw_teacher_id)
        except HTTPException:
            teacher_id = 0
        if teacher_id:
            cur.execute(
                "SELECT id_utilisateur FROM utilisateur WHERE id_utilisateur = %s AND role = 'enseignant'",
                (teacher_id,),
            )
            row = cur.fetchone()
            if row:
                return int(row["id_utilisateur"])

    raise HTTPException(status_code=404, detail="Enseignant introuvable dans PostgreSQL.")


def _resolve_or_create_class_id(cur, payload: dict) -> int:
    class_name = _require_text(payload.get("class_name") or payload.get("className"), "class_name")
    academic_year = str(payload.get("academic_year") or payload.get("academicYear") or "2025-2026").strip()
    level = str(payload.get("class_level") or payload.get("classLevel") or class_name).strip() or class_name

    cur.execute(
        """
        SELECT id_classe
        FROM classe
        WHERE lower(nom) = lower(%s) AND annee_scolaire = %s
        """,
        (class_name, academic_year),
    )
    row = cur.fetchone()
    if row:
        return int(row["id_classe"])

    cur.execute(
        """
        INSERT INTO classe (nom, niveau, annee_scolaire)
        VALUES (%s, %s, %s)
        RETURNING id_classe
        """,
        (class_name, level, academic_year),
    )
    return int(cur.fetchone()["id_classe"])


def _resolve_or_create_matiere_id(cur, payload: dict, class_id: int, teacher_id: int) -> int:
    subject_name = _require_text(payload.get("subject_name") or payload.get("subjectName"), "subject_name")
    coefficient = float(payload.get("subject_coefficient") or payload.get("subjectCoefficient") or 1)

    cur.execute(
        """
        INSERT INTO matiere (intitule, coefficient, id_classe, id_enseignant)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (intitule, id_classe)
        DO UPDATE SET coefficient = EXCLUDED.coefficient, id_enseignant = EXCLUDED.id_enseignant, updated_at = now()
        RETURNING id_matiere
        """,
        (subject_name, coefficient, class_id, teacher_id),
    )
    return int(cur.fetchone()["id_matiere"])


def _sync_web_qcm(payload: dict) -> dict:
    record = payload.get("qcm") if isinstance(payload.get("qcm"), dict) else payload
    if not isinstance(record, dict):
        raise HTTPException(status_code=400, detail="Payload QCM invalide.")

    qcm_code = _require_text(record.get("code"), "qcm.code")
    qcm_title = _require_text(record.get("title") or record.get("intitule"), "qcm.title")
    qcm_type = _qcm_type_from_web(record.get("audience") or record.get("type_qcm"))
    if qcm_type != "classe":
        raise HTTPException(status_code=400, detail="La synchronisation automatique est disponible pour les QCM de classe.")

    raw_questions = record.get("questions")
    if not isinstance(raw_questions, list) or not raw_questions:
        raise HTTPException(status_code=400, detail="Le QCM doit contenir au moins une question.")
    if len(raw_questions) > 60:
        raise HTTPException(status_code=400, detail="Le template actuel accepte 60 questions maximum.")
    questions = [_normalize_web_question(raw_question, index) for index, raw_question in enumerate(raw_questions)]

    sync_context = payload.get("context") if isinstance(payload.get("context"), dict) else {}
    context_payload = {
        **sync_context,
        "teacher_email": payload.get("teacher_email") or sync_context.get("teacher_email"),
        "teacher_id": record.get("teacherId") or record.get("teacher_id"),
        "class_name": payload.get("class_name") or sync_context.get("class_name"),
        "class_level": payload.get("class_level") or sync_context.get("class_level"),
        "academic_year": payload.get("academic_year") or sync_context.get("academic_year"),
        "subject_name": payload.get("subject_name") or sync_context.get("subject_name"),
        "subject_coefficient": payload.get("subject_coefficient") or sync_context.get("subject_coefficient"),
    }

    status = _qcm_status_from_web(record.get("status"))
    exam_date = str(record.get("examDate") or record.get("date_examen") or "").strip() or None
    published_at = record.get("publishedAt") or record.get("published_at")

    with connect() as conn:
        with conn.cursor() as cur:
            teacher_id = _resolve_teacher_id(cur, context_payload)
            class_id = _resolve_or_create_class_id(cur, context_payload)
            matiere_id = _resolve_or_create_matiere_id(cur, context_payload, class_id, teacher_id)

            cur.execute("SELECT id_qcm FROM qcm WHERE code = %s", (qcm_code,))
            existing = cur.fetchone()
            if existing:
                qcm_id = int(existing["id_qcm"])
                cur.execute(
                    """
                    UPDATE qcm
                    SET intitule = %s,
                        type_qcm = %s,
                        statut = %s,
                        date_examen = %s,
                        published_at = CASE WHEN %s = 'publie' THEN COALESCE(published_at, %s::timestamptz, now()) ELSE published_at END,
                        id_createur = %s,
                        id_matiere = %s,
                        id_classe = %s,
                        updated_at = now()
                    WHERE id_qcm = %s
                    """,
                    (qcm_title, qcm_type, status, exam_date, status, published_at, teacher_id, matiere_id, class_id, qcm_id),
                )
                cur.execute("DELETE FROM question WHERE id_qcm = %s", (qcm_id,))
            else:
                cur.execute(
                    """
                    INSERT INTO qcm (
                        code,
                        intitule,
                        type_qcm,
                        statut,
                        date_examen,
                        published_at,
                        id_createur,
                        id_matiere,
                        id_classe
                    )
                    VALUES (%s, %s, %s, %s, %s, CASE WHEN %s = 'publie' THEN COALESCE(%s::timestamptz, now()) ELSE NULL END, %s, %s, %s)
                    RETURNING id_qcm
                    """,
                    (qcm_code, qcm_title, qcm_type, status, exam_date, status, published_at, teacher_id, matiere_id, class_id),
                )
                qcm_id = int(cur.fetchone()["id_qcm"])

            for question in questions:
                cur.execute(
                    """
                    INSERT INTO question (intitule, numero, ponderation, nbr_reponse, type_question, id_qcm)
                    VALUES (%s, %s, %s, %s, %s, %s)
                    RETURNING id_question
                    """,
                    (
                        question["prompt"],
                        question["numero"],
                        question["points"],
                        len(question["options"]),
                        question["type_question"],
                        qcm_id,
                    ),
                )
                question_id = int(cur.fetchone()["id_question"])
                for option in question["options"]:
                    cur.execute(
                        """
                        INSERT INTO reponse (label, intitule, est_correcte, id_question)
                        VALUES (%s, %s, %s, %s)
                        """,
                        (option["label"], option["text"], option["is_correct"], question_id),
                    )

            session_id = None
            if status == "publie":
                cur.execute(
                    """
                    SELECT id_session
                    FROM session_correction
                    WHERE id_qcm = %s
                    ORDER BY id_session
                    LIMIT 1
                    """,
                    (qcm_id,),
                )
                existing_session = cur.fetchone()
                if existing_session:
                    session_id = int(existing_session["id_session"])
                    cur.execute(
                        """
                        UPDATE session_correction
                        SET libelle = %s,
                            statut = CASE WHEN statut = 'terminee' THEN statut ELSE 'planifiee'::statut_session END,
                            id_enseignant = %s,
                            id_classe = %s,
                            id_concours = NULL,
                            updated_at = now()
                        WHERE id_session = %s
                        """,
                        (qcm_title, teacher_id, class_id, session_id),
                    )
                else:
                    cur.execute(
                        """
                        INSERT INTO session_correction (libelle, statut, email_status, id_qcm, id_enseignant, id_classe)
                        VALUES (%s, 'planifiee', 'en_attente', %s, %s, %s)
                        RETURNING id_session
                        """,
                        (qcm_title, qcm_id, teacher_id, class_id),
                    )
                    session_id = int(cur.fetchone()["id_session"])
        conn.commit()

    return {
        "status": "ok",
        "qcm_id": f"qcm-{qcm_id}",
        "session_id": f"session-{session_id}" if session_id else None,
        "qcm_code": qcm_code,
        "question_count": len(questions),
        "max_score": sum(float(question["points"]) for question in questions),
    }


def _save_web_qcm_draft(payload: dict) -> dict:
    record = payload.get("qcm") if isinstance(payload.get("qcm"), dict) else payload
    if not isinstance(record, dict):
        raise HTTPException(status_code=400, detail="Payload brouillon invalide.")

    sync_context = payload.get("context") if isinstance(payload.get("context"), dict) else {}
    context_payload = {
        **sync_context,
        "teacher_email": payload.get("teacher_email") or sync_context.get("teacher_email"),
        "teacher_id": record.get("teacherId") or record.get("teacher_id"),
    }
    title = str(record.get("title") or record.get("intitule") or "Brouillon QCM").strip() or "Brouillon QCM"

    with connect() as conn:
        with conn.cursor() as cur:
            teacher_id = _resolve_teacher_id(cur, context_payload)
            qcm_id = None
            qcm_code = str(record.get("code") or "").strip()
            if qcm_code:
                cur.execute("SELECT id_qcm FROM qcm WHERE code = %s", (qcm_code,))
                row = cur.fetchone()
                if row:
                    qcm_id = int(row["id_qcm"])

            if qcm_id:
                cur.execute("SELECT id_brouillon FROM brouillon_autosave WHERE id_qcm = %s", (qcm_id,))
                existing = cur.fetchone()
            else:
                existing = None
            if existing:
                cur.execute(
                    """
                    UPDATE brouillon_autosave
                    SET intitule = %s,
                        payload = %s::jsonb,
                        id_createur = %s,
                        date_maj = now(),
                        updated_at = now()
                    WHERE id_brouillon = %s
                    RETURNING id_brouillon
                    """,
                    (title, json.dumps(payload, ensure_ascii=False), teacher_id, existing["id_brouillon"]),
                )
            else:
                cur.execute(
                    """
                    INSERT INTO brouillon_autosave (intitule, payload, id_createur, id_qcm, date_maj)
                    VALUES (%s, %s::jsonb, %s, %s, now())
                    RETURNING id_brouillon
                    """,
                    (title, json.dumps(payload, ensure_ascii=False), teacher_id, qcm_id),
                )
            draft_id = int(cur.fetchone()["id_brouillon"])
        conn.commit()

    return {"status": "ok", "draft_id": f"draft-{draft_id}", "qcm_id": f"qcm-{qcm_id}" if qcm_id else None}


def _fetch_sessions_overview(viewer: dict | None = None) -> list[dict]:
    params: list[object] = []
    where_clause = ""
    if _is_teacher_viewer(viewer):
        where_clause = """
        WHERE (
            s.id_enseignant = %s
            OR q.id_createur = %s
            OR m.id_enseignant = %s
        )
        """
        params.extend([viewer["id"], viewer["id"], viewer["id"]])

    query = f"""
        SELECT
            s.id_session,
            s.libelle,
            s.statut::text AS statut,
            s.email_status::text AS email_status,
            s.id_qcm,
            q.code AS qcm_code,
            q.intitule AS qcm_title,
            q.type_qcm::text AS qcm_type,
            q.statut::text AS qcm_status,
            q.date_examen,
            q.updated_at AS qcm_updated_at,
            s.id_enseignant,
            u.email AS teacher_email,
            u.prenom AS teacher_first_name,
            u.nom AS teacher_last_name,
            s.id_classe,
            c.nom AS class_name,
            m.intitule AS subject_name,
            COALESCE(q_totals.question_count, 0)::int AS question_count,
            COALESCE(q_totals.max_score, 0)::float AS max_score,
            COUNT(i.id_instance)::int AS copies,
            COALESCE(ROUND(AVG(i.score)::numeric, 2), 0)::float AS average_score,
            COALESCE(SUM(CASE WHEN i.statut_scan = 'a_verifier' THEN 1 ELSE 0 END), 0)::int AS manual_checks
        FROM session_correction s
        JOIN qcm q ON q.id_qcm = s.id_qcm
        JOIN utilisateur u ON u.id_utilisateur = s.id_enseignant
        JOIN matiere m ON m.id_matiere = q.id_matiere
        LEFT JOIN classe c ON c.id_classe = s.id_classe
        LEFT JOIN instance_qcm i ON i.id_session = s.id_session
        LEFT JOIN (
            SELECT id_qcm, COUNT(*) AS question_count, COALESCE(SUM(ponderation), 0) AS max_score
            FROM question
            GROUP BY id_qcm
        ) q_totals ON q_totals.id_qcm = q.id_qcm
        {where_clause}
        GROUP BY s.id_session, q.id_qcm, u.id_utilisateur, c.id_classe, m.id_matiere, q_totals.question_count, q_totals.max_score
        ORDER BY s.id_session DESC
    """
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(query, params)
            rows = cur.fetchall()
    return [
        {
            "id": f"session-{int(row['id_session'])}",
            "qcm_id": f"qcm-{int(row['id_qcm'])}",
            "qcm_code": row["qcm_code"],
            "qcm_status": row["qcm_status"],
            "qcm_type": row["qcm_type"],
            "question_count": int(row["question_count"] or 0),
            "max_score": float(row["max_score"] or 0),
            "subject_name": row.get("subject_name"),
            "label": row["libelle"],
            "exam_date": row.get("date_examen").isoformat() if row.get("date_examen") else "",
            "updated_at": _as_iso(row.get("qcm_updated_at")),
            "state": row["statut"],
            "email_status": row["email_status"],
            "teacher": {
                "id": int(row["id_enseignant"]),
                "email": row.get("teacher_email"),
                "first_name": row.get("teacher_first_name"),
                "last_name": row.get("teacher_last_name"),
            },
            "class": {
                "id": int(row["id_classe"]) if row.get("id_classe") else None,
                "name": row.get("class_name"),
            },
            "copies": int(row["copies"]),
            "average_score": float(row["average_score"]),
            "manual_checks": int(row["manual_checks"]),
        }
        for row in rows
    ]


def _fetch_classes_with_students(viewer: dict | None = None) -> list[dict]:
    params: list[object] = []
    class_scope = ""
    subject_scope = ""
    if _is_teacher_viewer(viewer):
        class_scope = """
                WHERE EXISTS (
                    SELECT 1
                    FROM matiere m
                    WHERE m.id_classe = c.id_classe
                      AND m.id_enseignant = %s
                )
        """
        subject_scope = "WHERE m.id_enseignant = %s"
        params.append(viewer["id"])

    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT id_classe, nom, niveau, annee_scolaire
                FROM classe c
                {class_scope}
                ORDER BY id_classe
                """,
                params,
            )
            classes = cur.fetchall()
            cur.execute(
                """
                SELECT id_etudiant, numero_etudiant, nom, prenom, email, id_classe, actif
                FROM etudiant
                ORDER BY id_classe, numero_etudiant
                """
            )
            students = cur.fetchall()
            subject_params = [viewer["id"]] if _is_teacher_viewer(viewer) else []
            cur.execute(
                f"""
                SELECT
                    m.id_matiere,
                    m.intitule,
                    m.coefficient,
                    m.id_classe,
                    m.id_enseignant,
                    u.prenom AS teacher_first_name,
                    u.nom AS teacher_last_name,
                    u.email AS teacher_email
                FROM matiere m
                JOIN utilisateur u ON u.id_utilisateur = m.id_enseignant
                {subject_scope}
                ORDER BY m.id_classe, m.intitule
                """,
                subject_params,
            )
            subjects = cur.fetchall()

    subjects_by_class: dict[int, list[dict]] = {}
    for subject in subjects:
        class_id = int(subject["id_classe"])
        subjects_by_class.setdefault(class_id, []).append(
            {
                "id": int(subject["id_matiere"]),
                "name": subject["intitule"],
                "coefficient": float(subject["coefficient"] or 1),
                "teacher_id": int(subject["id_enseignant"]),
                "teacher_name": " ".join(
                    str(part).strip()
                    for part in (subject.get("teacher_first_name"), subject.get("teacher_last_name"))
                    if part
                ).strip(),
                "teacher_email": subject.get("teacher_email"),
            }
        )

    students_by_class: dict[int, list[dict]] = {}
    for student in students:
        class_id = int(student["id_classe"])
        students_by_class.setdefault(class_id, []).append(
            {
                "id": int(student["id_etudiant"]),
                "number": student["numero_etudiant"],
                "first_name": student["prenom"],
                "last_name": student["nom"],
                "email": student["email"],
                "active": bool(student["actif"]),
            }
        )
    return [
        {
            "id": int(row["id_classe"]),
            "name": row["nom"],
            "level": row["niveau"],
            "academic_year": row["annee_scolaire"],
            "subjects": subjects_by_class.get(int(row["id_classe"]), []),
            "students": students_by_class.get(int(row["id_classe"]), []),
        }
        for row in classes
    ]



def _require_admin_viewer(viewer: dict | None) -> dict:
    if not viewer or viewer.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Action reservee a l'administrateur.")
    return viewer


def _normalize_academic_year(value: object) -> str:
    text = str(value or "2025-2026").strip()
    return text or "2025-2026"


def _normalize_student_payload(payload: dict, class_id: int) -> dict:
    first_name = _require_text(payload.get("first_name") or payload.get("prenom"), "prenom")
    last_name = _require_text(payload.get("last_name") or payload.get("nom"), "nom")
    email_value = _require_text(payload.get("email"), "email").lower()
    number = _normalize_student_number(payload.get("number") or payload.get("numero_etudiant") or payload.get("student_number"))
    if not re.fullmatch(r"\d{6}", number):
        raise HTTPException(status_code=400, detail="Le numero etudiant doit contenir exactement 6 chiffres.")
    return {
        "first_name": first_name,
        "last_name": last_name,
        "email": email_value,
        "number": number,
        "class_id": class_id,
        "source_import": str(payload.get("source_import") or "manuel").strip() or "manuel",
    }


def _insert_student(cur, payload: dict, class_id: int) -> int:
    student = _normalize_student_payload(payload, class_id)
    cur.execute("SELECT id_classe FROM classe WHERE id_classe = %s", (class_id,))
    if not cur.fetchone():
        raise HTTPException(status_code=404, detail="Classe introuvable.")
    cur.execute(
        """
        INSERT INTO personne (nom, prenom, email, id_groupe, type_personne, numero_etudiant)
        VALUES (%s, %s, %s, %s, 'etudiant', %s)
        RETURNING id_personne
        """,
        (student["last_name"], student["first_name"], student["email"], class_id, student["number"]),
    )
    person_id = int(cur.fetchone()["id_personne"])
    cur.execute(
        """
        INSERT INTO etudiant (id_personne, numero_etudiant, nom, prenom, email, id_classe, source_import, actif)
        VALUES (%s, %s, %s, %s, %s, %s, %s, TRUE)
        RETURNING id_etudiant
        """,
        (
            person_id,
            student["number"],
            student["last_name"],
            student["first_name"],
            student["email"],
            class_id,
            student["source_import"],
        ),
    )
    return int(cur.fetchone()["id_etudiant"])


def _link_pending_instances_for_student(cur, student_id: int) -> int:
    cur.execute(
        """
        UPDATE instance_qcm i
        SET id_etudiant = e.id_etudiant,
            id_personne = e.id_personne,
            updated_at = now()
        FROM etudiant e, session_correction s
        WHERE e.id_etudiant = %s
          AND e.actif = TRUE
          AND s.id_classe = e.id_classe
          AND i.id_session = s.id_session
          AND i.id_etudiant IS NULL
          AND i.id_candidat IS NULL
          AND i.numero_etudiant_detecte = e.numero_etudiant
        """,
        (student_id,),
    )
    return int(cur.rowcount or 0)


def _reconcile_pending_student_associations(session_id: int) -> int:
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE instance_qcm i
                SET id_etudiant = e.id_etudiant,
                    id_personne = e.id_personne,
                    updated_at = now()
                FROM etudiant e, session_correction s
                WHERE i.id_session = %s
                  AND s.id_session = i.id_session
                  AND e.id_classe = s.id_classe
                  AND e.actif = TRUE
                  AND i.id_etudiant IS NULL
                  AND i.id_candidat IS NULL
                  AND i.numero_etudiant_detecte = e.numero_etudiant
                """,
                (session_id,),
            )
            linked = int(cur.rowcount or 0)
        conn.commit()
    return linked


def _csv_rows_from_upload(file_bytes: bytes, filename: str) -> list[dict]:
    lower_name = filename.lower()
    if lower_name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ModuleNotFoundError as exc:
            raise HTTPException(status_code=400, detail="Import Excel indisponible: installez openpyxl ou utilisez un CSV.") from exc
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip().lower() for value in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
            if any(value not in (None, "") for value in row)
        ]

    sample = file_bytes.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(sample[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.DictReader(io.StringIO(sample), dialect=dialect))


def _student_payload_from_import_row(row: dict) -> dict:
    normalized = {str(k or "").strip().lower(): v for k, v in row.items()}
    return {
        "number": normalized.get("numero_etudiant") or normalized.get("identifiant") or normalized.get("student_number") or normalized.get("number") or normalized.get("id"),
        "first_name": normalized.get("prenom") or normalized.get("first_name") or normalized.get("firstname"),
        "last_name": normalized.get("nom") or normalized.get("last_name") or normalized.get("lastname"),
        "email": normalized.get("email") or normalized.get("mail") or normalized.get("adresse email"),
        "source_import": "import",
    }


def _fetch_professors() -> list[dict]:
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT
                    u.id_utilisateur,
                    u.prenom,
                    u.nom,
                    u.email,
                    u.actif,
                    COALESCE(jsonb_agg(
                        DISTINCT jsonb_build_object(
                            'id', m.id_matiere,
                            'name', m.intitule,
                            'coefficient', m.coefficient,
                            'class_id', c.id_classe,
                            'class_name', c.nom
                        )
                    ) FILTER (WHERE m.id_matiere IS NOT NULL), '[]'::jsonb) AS subjects,
                    COUNT(DISTINCT q.id_qcm)::int AS qcm_count,
                    COUNT(DISTINCT s.id_session)::int AS session_count
                FROM utilisateur u
                LEFT JOIN matiere m ON m.id_enseignant = u.id_utilisateur
                LEFT JOIN classe c ON c.id_classe = m.id_classe
                LEFT JOIN qcm q ON q.id_createur = u.id_utilisateur OR q.id_matiere = m.id_matiere
                LEFT JOIN session_correction s ON s.id_enseignant = u.id_utilisateur OR s.id_qcm = q.id_qcm
                WHERE u.role = 'enseignant'
                  AND u.actif = TRUE
                GROUP BY u.id_utilisateur, u.prenom, u.nom, u.email, u.actif
                ORDER BY u.id_utilisateur
                """
            )
            rows = cur.fetchall()
    return [
        {
            "id": int(row["id_utilisateur"]),
            "first_name": row["prenom"],
            "last_name": row["nom"],
            "email": row["email"],
            "active": bool(row["actif"]),
            "subjects": row["subjects"] or [],
            "qcm_count": int(row["qcm_count"] or 0),
            "session_count": int(row["session_count"] or 0),
        }
        for row in rows
    ]


def _normalize_professor_payload(payload: dict, existing_password: str | None = None) -> dict:
    first_name = _require_text(payload.get("first_name") or payload.get("prenom"), "prenom")
    last_name = _require_text(payload.get("last_name") or payload.get("nom"), "nom")
    email = _require_text(payload.get("email"), "email").lower()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise HTTPException(status_code=400, detail="Email professeur invalide.")
    password = str(payload.get("password") or payload.get("mot_de_passe") or existing_password or os.environ.get("DEFAULT_PROFESSOR_PASSWORD", "")).strip()
    if len(password) < 6:
        raise HTTPException(status_code=400, detail="Le mot de passe professeur doit contenir au moins 6 caracteres.")
    subject_name = _require_text(payload.get("subject_name") or payload.get("matiere"), "matiere")
    try:
        class_id = int(payload.get("class_id") or payload.get("id_classe") or 0)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Classe invalide.") from exc
    if not class_id:
        raise HTTPException(status_code=400, detail="Classe obligatoire pour affecter le professeur.")
    try:
        coefficient = float(payload.get("coefficient") or 1)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Coefficient matiere invalide.") from exc
    if coefficient <= 0:
        raise HTTPException(status_code=400, detail="Le coefficient doit etre positif.")
    return {
        "first_name": first_name,
        "last_name": last_name,
        "email": email,
        "password": password,
        "subject_name": subject_name,
        "class_id": class_id,
        "coefficient": coefficient,
        "active": bool(payload.get("active", payload.get("actif", True))),
    }


def _upsert_professor_subject(cur, professor_id: int, payload: dict, subject_id: int | None = None) -> int:
    cur.execute("SELECT id_classe FROM classe WHERE id_classe = %s", (payload["class_id"],))
    if not cur.fetchone():
        raise HTTPException(status_code=404, detail="Classe introuvable pour l'affectation du professeur.")

    if subject_id:
        cur.execute(
            """
            UPDATE matiere
            SET intitule = %s, coefficient = %s, id_classe = %s, id_enseignant = %s, updated_at = now()
            WHERE id_matiere = %s AND id_enseignant = %s
            RETURNING id_matiere
            """,
            (payload["subject_name"], payload["coefficient"], payload["class_id"], professor_id, subject_id, professor_id),
        )
        row = cur.fetchone()
        if row:
            return int(row["id_matiere"])

    cur.execute(
        """
        INSERT INTO matiere (intitule, coefficient, id_classe, id_enseignant)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (intitule, id_classe)
        DO UPDATE SET coefficient = EXCLUDED.coefficient, id_enseignant = EXCLUDED.id_enseignant, updated_at = now()
        RETURNING id_matiere
        """,
        (payload["subject_name"], payload["coefficient"], payload["class_id"], professor_id),
    )
    return int(cur.fetchone()["id_matiere"])


def _compute_weighted_score(
    predicted: dict[str, list[str]],
    questions: list[dict],
) -> tuple[float, float, list[int], list[int], dict[str, list[str]]]:
    total_points = 0.0
    earned_points = 0.0
    correct: list[int] = []
    wrong: list[int] = []
    expected_map: dict[str, list[str]] = {}
    for question in questions:
        numero = int(question["numero"])
        points = float(question["ponderation"])
        expected = list(question["expected"])
        got = predicted.get(str(numero), [])
        total_points += points
        expected_map[str(numero)] = expected
        if set(got) == set(expected):
            earned_points += points
            correct.append(numero)
        else:
            wrong.append(numero)
    return earned_points, total_points, correct, wrong, expected_map


def _normalize_detected_answer_values(raw_answers: dict) -> dict[str, list[str]]:
    valid_labels = {"A", "B", "C", "D"}
    normalized: dict[str, list[str]] = {}
    for raw_key, raw_value in (raw_answers or {}).items():
        key = str(raw_key).strip()
        if not key.isdigit():
            continue
        values = raw_value if isinstance(raw_value, (list, tuple, set)) else [raw_value]
        choices: list[str] = []
        for value in values:
            labels = re.findall(r"[A-D]", str(value).strip().upper())
            for label in labels:
                if label in valid_labels and label not in choices:
                    choices.append(label)
        normalized[key] = [label for label in ("A", "B", "C", "D") if label in choices]
    return normalized


def _normalize_detected_answers(raw_answers: dict, questions: list[dict]) -> dict[str, list[str]]:
    question_numbers = {str(question["numero"]) for question in questions}
    normalized = _normalize_detected_answer_values(raw_answers)
    return {key: value for key, value in normalized.items() if key in question_numbers}


def _extract_answers_from_detection_payload(payload: dict) -> dict:
    raw_answers = payload.get("answers")
    if isinstance(raw_answers, dict):
        return raw_answers

    extracted: dict[str, list[str] | str] = {}
    for key, value in payload.items():
        match = re.fullmatch(r"[qQ](\d+)", str(key).strip())
        if not match:
            continue
        question_number = match.group(1)
        extracted[question_number] = value
    return extracted


def _require_ai_results_key(x_api_key: str | None) -> None:
    if AI_RESULTS_API_KEY and x_api_key != AI_RESULTS_API_KEY:
        raise HTTPException(status_code=401, detail="Cle API IA invalide ou manquante.")


def _as_list_payload(value: object) -> list[object]:
    if isinstance(value, list):
        return value
    return [value]


def _normalize_answers_payload(raw: object) -> dict:
    if isinstance(raw, dict):
        return raw
    if not isinstance(raw, list):
        return {}

    answers: dict[str, object] = {}
    for item in raw:
        if not isinstance(item, dict):
            continue
        q = item.get("question") or item.get("numero") or item.get("number") or item.get("q")
        value = (
            item.get("answer")
            or item.get("answers")
            or item.get("reponse")
            or item.get("reponses")
            or item.get("choices")
            or item.get("selected")
        )
        if q is None or value is None:
            continue
        answers[str(q)] = value
    return answers


def _normalize_ai_result_payload(payload: dict, inherited_session_id: object | None = None) -> dict:
    normalized = dict(payload)
    student = normalized.get("student") if isinstance(normalized.get("student"), dict) else {}
    candidate = normalized.get("candidate") if isinstance(normalized.get("candidate"), dict) else {}

    student_number = (
        normalized.get("student_id_detected")
        or normalized.get("numero_etudiant_detecte")
        or normalized.get("student_number")
        or normalized.get("student_id")
        or normalized.get("studentId")
        or normalized.get("numero_etudiant")
        or normalized.get("id")
        or student.get("student_id")
        or student.get("studentId")
        or student.get("numero_etudiant")
        or student.get("number")
        or student.get("id")
    )
    if student_number is not None:
        normalized["student_id_detected"] = student_number

    candidate_number = (
        normalized.get("candidate_id_detected")
        or normalized.get("matricule_candidat_detecte")
        or normalized.get("candidate_number")
        or normalized.get("candidate_id")
        or normalized.get("candidateId")
        or normalized.get("matricule")
        or candidate.get("candidate_id")
        or candidate.get("candidateId")
        or candidate.get("matricule")
        or candidate.get("number")
        or candidate.get("id")
    )
    if candidate_number is not None:
        normalized["candidate_id_detected"] = candidate_number

    answers = (
        normalized.get("answers")
        or normalized.get("detected_answers")
        or normalized.get("reponses")
        or normalized.get("responses")
        or normalized.get("questions")
        or normalized.get("results")
    )
    normalized_answers = _normalize_answers_payload(answers)
    if normalized_answers:
        normalized["answers"] = normalized_answers

    session_id = normalized.get("session_id") or normalized.get("sessionId") or inherited_session_id
    if session_id is not None:
        normalized["session_id"] = session_id

    if "answers_method" not in normalized:
        normalized["answers_method"] = normalized.get("method") or normalized.get("methode") or "tensorflow"
    return normalized


def _normalize_name_for_match(value: object) -> str:
    text = str(value or "").strip().lower()
    text = "".join(
        ch for ch in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(ch)
    )
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _extract_detected_full_name(payload: dict) -> str:
    student = payload.get("student") if isinstance(payload.get("student"), dict) else {}
    parts = [
        payload.get("nom_complet")
        or payload.get("full_name")
        or payload.get("name")
        or student.get("nom_complet")
        or student.get("full_name")
        or student.get("name")
    ]
    if not parts[0]:
        parts = [
            payload.get("prenom") or payload.get("first_name") or student.get("prenom") or student.get("first_name"),
            payload.get("nom") or payload.get("last_name") or student.get("nom") or student.get("last_name"),
        ]
    return " ".join(str(part).strip() for part in parts if part).strip()


def _detected_name_matches_db(detected_name: str, first_name: object, last_name: object) -> bool:
    detected = _normalize_name_for_match(detected_name)
    if not detected:
        return True
    first = _normalize_name_for_match(first_name)
    last = _normalize_name_for_match(last_name)
    full_1 = f"{first} {last}".strip()
    full_2 = f"{last} {first}".strip()
    if detected in {first, last, full_1, full_2}:
        return True
    detected_tokens = set(detected.split())
    db_tokens = set(full_1.split())
    return bool(detected_tokens) and detected_tokens.issubset(db_tokens)


def _match_detected_participant(session_ctx: dict, payload: dict) -> dict:
    qcm_type = session_ctx.get("qcm_type")
    detected_name = _extract_detected_full_name(payload)
    raw_student = _normalize_student_number(
        payload.get("student_id_detected")
        or payload.get("numero_etudiant_detecte")
        or payload.get("student_number")
        or payload.get("id")
        or ""
    )
    raw_candidate = str(
        payload.get("candidate_id_detected")
        or payload.get("matricule_candidat_detecte")
        or payload.get("candidate_number")
        or ""
    ).strip()

    with connect() as conn:
        with conn.cursor() as cur:
            if qcm_type == "classe" and raw_student:
                cur.execute(
                    """
                    SELECT id_etudiant, numero_etudiant, nom, prenom, email
                    FROM etudiant
                    WHERE numero_etudiant = %s
                      AND id_classe = %s
                      AND actif = TRUE
                    """,
                    (raw_student, session_ctx["id_classe"]),
                )
                row = cur.fetchone()
                if row:
                    matched = {
                        "type": "student",
                        "id": int(row["id_etudiant"]),
                        "number": row["numero_etudiant"],
                        "first_name": row["prenom"],
                        "last_name": row["nom"],
                        "email": row["email"],
                    }
                    if detected_name and not _detected_name_matches_db(detected_name, row["prenom"], row["nom"]):
                        return {
                            "kind": "student",
                            "id_etudiant": int(row["id_etudiant"]),
                            "id_candidat": None,
                            "detected_student": raw_student,
                            "detected_candidate": None,
                            "detected_name": detected_name,
                            "matched": matched,
                            "match_warning": (
                                f"Nom detecte '{detected_name}' incompatible avec l'etudiant "
                                f"{row['prenom']} {row['nom']} pour le numero {raw_student}."
                            ),
                        }
                    return {
                        "kind": "student",
                        "id_etudiant": int(row["id_etudiant"]),
                        "id_candidat": None,
                        "detected_student": raw_student,
                        "detected_candidate": None,
                        "detected_name": detected_name or None,
                        "matched": matched,
                    }
            if qcm_type == "concours" and raw_candidate:
                cur.execute(
                    """
                    SELECT ca.id_candidat, ca.matricule, ca.nom, ca.prenom, ca.email
                    FROM candidat ca
                    JOIN inscription_concours ic ON ic.id_candidat = ca.id_candidat
                    WHERE ca.matricule = %s
                      AND ic.id_concours = %s
                    """,
                    (raw_candidate, session_ctx["id_concours"]),
                )
                row = cur.fetchone()
                if row:
                    matched = {
                        "type": "candidate",
                        "id": int(row["id_candidat"]),
                        "number": row["matricule"],
                        "first_name": row["prenom"],
                        "last_name": row["nom"],
                        "email": row["email"],
                    }
                    if detected_name and not _detected_name_matches_db(detected_name, row["prenom"], row["nom"]):
                        return {
                            "kind": "candidate",
                            "id_etudiant": None,
                            "id_candidat": int(row["id_candidat"]),
                            "detected_student": None,
                            "detected_candidate": raw_candidate,
                            "detected_name": detected_name,
                            "matched": matched,
                            "match_warning": (
                                f"Nom detecte '{detected_name}' incompatible avec le candidat "
                                f"{row['prenom']} {row['nom']} pour le matricule {raw_candidate}."
                            ),
                        }
                    return {
                        "kind": "candidate",
                        "id_etudiant": None,
                        "id_candidat": int(row["id_candidat"]),
                        "detected_student": None,
                        "detected_candidate": raw_candidate,
                        "detected_name": detected_name or None,
                        "matched": matched,
                    }

    return {
        "kind": None,
        "id_etudiant": None,
        "id_candidat": None,
        "detected_student": raw_student or None,
        "detected_candidate": raw_candidate or None,
        "detected_name": detected_name or None,
        "matched": None,
    }


def _extract_detected_student_number(payload: dict) -> str:
    return _normalize_student_number(
        payload.get("student_id_detected")
        or payload.get("numero_etudiant_detecte")
        or payload.get("student_number")
        or payload.get("id")
        or ""
    )


def _resolve_session_from_detection_payload(payload: dict) -> dict:
    """Find the correction session from the detected student id when no session is provided."""
    detected_student = _extract_detected_student_number(payload)
    if not detected_student:
        raise HTTPException(
            status_code=400,
            detail="Impossible de detecter la session: champ id/numero_etudiant manquant dans le JSON IA.",
        )

    query = """
        SELECT
            s.id_session,
            s.libelle,
            s.statut::text AS statut,
            s.email_status::text AS email_status,
            s.id_qcm,
            q.code AS qcm_code,
            q.intitule AS qcm_title,
            q.type_qcm::text AS qcm_type,
            q.statut::text AS qcm_status,
            q.date_examen,
            s.id_enseignant,
            u.nom AS enseignant_nom,
            u.prenom AS enseignant_prenom,
            s.id_classe,
            c.nom AS classe_nom,
            c.niveau AS classe_niveau,
            s.id_concours,
            co.intitule AS concours_intitule
        FROM etudiant e
        JOIN session_correction s ON s.id_classe = e.id_classe
        JOIN qcm q ON q.id_qcm = s.id_qcm
        JOIN utilisateur u ON u.id_utilisateur = s.id_enseignant
        LEFT JOIN classe c ON c.id_classe = s.id_classe
        LEFT JOIN concours co ON co.id_concours = s.id_concours
        WHERE e.numero_etudiant = %s
          AND e.actif = TRUE
          AND q.type_qcm = 'classe'
          AND q.statut = 'publie'
        ORDER BY
          CASE s.statut
            WHEN 'en_cours' THEN 1
            WHEN 'planifiee' THEN 2
            WHEN 'terminee' THEN 3
            ELSE 4
          END,
          s.started_at DESC NULLS LAST,
          s.id_session DESC
    """
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(query, (detected_student,))
            rows = cur.fetchall()

    if not rows:
        raise HTTPException(
            status_code=404,
            detail=f"Aucune session/QCM publie trouve pour l'etudiant detecte {detected_student}.",
        )

    active_rows = [row for row in rows if row["statut"] in {"en_cours", "planifiee"}]
    candidates = active_rows or rows
    if len(candidates) > 1:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "ambiguous_session",
                "message": "Plusieurs sessions possibles pour cet etudiant. Envoyez plutot vers /api/sessions/{session_id}/detected-result.",
                "detected_student": detected_student,
                "sessions": [
                    {
                        "session_id": f"session-{int(row['id_session'])}",
                        "id_session": int(row["id_session"]),
                        "qcm_id": int(row["id_qcm"]),
                        "qcm_code": row["qcm_code"],
                        "qcm_title": row["qcm_title"],
                        "status": row["statut"],
                        "class_name": row.get("classe_nom"),
                    }
                    for row in candidates
                ],
            },
        )

    return candidates[0]


def _next_attempt_no(cur, session_id: int, id_etudiant: int | None, id_candidat: int | None) -> int:
    if id_etudiant is not None:
        cur.execute(
            "SELECT COALESCE(MAX(tentative_no), 0) + 1 AS next_no FROM instance_qcm WHERE id_session = %s AND id_etudiant = %s",
            (session_id, id_etudiant),
        )
        return int(cur.fetchone()["next_no"])
    if id_candidat is not None:
        cur.execute(
            "SELECT COALESCE(MAX(tentative_no), 0) + 1 AS next_no FROM instance_qcm WHERE id_session = %s AND id_candidat = %s",
            (session_id, id_candidat),
        )
        return int(cur.fetchone()["next_no"])
    return 1


def _save_scan_image(payload: bytes, filename: str | None, session_id: int) -> Path:
    ext = (Path(filename or "").suffix or ".jpg").lower()
    if ext not in (".jpg", ".jpeg", ".png"):
        ext = ".jpg"
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = SUBMISSION_DIR / f"session-{session_id}-{stamp}{ext}"
    target.write_bytes(payload)
    return target


def _store_scan_result(
    session_ctx: dict,
    image_path: Path,
    answers: dict[str, list[str]],
    score: float,
    max_score: float,
    correct_questions: list[int],
    questions: list[dict],
    method_used: str,
) -> int:
    overlay_value = str(DEBUG_OVERLAY_PATH) if DEBUG_OVERLAY_PATH.exists() else None
    score_by_num = {
        int(question["numero"]): float(question["ponderation"]) if int(question["numero"]) in correct_questions else 0.0
        for question in questions
    }
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO instance_qcm (
                    id_session,
                    id_qcm,
                    image_path,
                    overlay_path,
                    raw_answers,
                    score,
                    max_score,
                    methode_reponses,
                    methode_id,
                    modele_reponses_version,
                    raw_id_detection,
                    statut_scan
                )
                VALUES (%s, %s, %s, %s, %s::jsonb, %s, %s, %s, %s, %s, %s::jsonb, %s)
                RETURNING id_instance
                """,
                (
                    session_ctx["id_session"],
                    session_ctx["id_qcm"],
                    str(image_path),
                    overlay_value,
                    json.dumps(answers, ensure_ascii=False),
                    score,
                    max_score,
                    method_used,
                    "pending_id_model",
                    MODEL_PATH.name,
                    json.dumps({"status": "id_model_not_connected_yet"}, ensure_ascii=False),
                    "a_verifier",
                ),
            )
            instance_id = int(cur.fetchone()["id_instance"])
            for question in questions:
                numero = int(question["numero"])
                cur.execute(
                    """
                    INSERT INTO reponse_soumise (
                        id_instance,
                        id_question,
                        est_correcte,
                        score_obtenu,
                        confidence
                    )
                    VALUES (%s, %s, %s, %s, %s)
                    RETURNING id_reponse_soumise
                    """,
                    (
                        instance_id,
                        question["id_question"],
                        numero in correct_questions,
                        score_by_num.get(numero, 0.0),
                        None,
                    ),
                )
                submission_answer_id = int(cur.fetchone()["id_reponse_soumise"])
                for label in answers.get(str(numero), []):
                    choice = question["choices"].get(label)
                    if not choice:
                        continue
                    cur.execute(
                        """
                        INSERT INTO reponse_soumise_choix (id_reponse_soumise, id_reponse)
                        VALUES (%s, %s)
                        ON CONFLICT DO NOTHING
                        """,
                        (submission_answer_id, choice["id_reponse"]),
                    )
        conn.commit()
    return instance_id


def _store_detected_result(
    session_ctx: dict,
    payload: dict,
    participant: dict,
    answers: dict[str, list[str]],
    score: float,
    max_score: float,
    correct_questions: list[int],
    questions: list[dict],
    method_used: str,
) -> int:
    score_by_num = {
        int(question["numero"]): float(question["ponderation"]) if int(question["numero"]) in correct_questions else 0.0
        for question in questions
    }
    model_versions = payload.get("model_versions") if isinstance(payload.get("model_versions"), dict) else {}
    raw_id_detection = dict(payload.get("raw_id_detection")) if isinstance(payload.get("raw_id_detection"), dict) else {}
    if participant.get("detected_name"):
        raw_id_detection["detected_name"] = participant.get("detected_name")
    if participant.get("match_warning"):
        raw_id_detection["match_warning"] = participant.get("match_warning")
    status = "valide" if participant.get("matched") and not participant.get("match_warning") else "a_verifier"

    with connect() as conn:
        with conn.cursor() as cur:
            tentative_no = _next_attempt_no(
                cur,
                int(session_ctx["id_session"]),
                participant.get("id_etudiant"),
                participant.get("id_candidat"),
            )
            cur.execute(
                """
                INSERT INTO instance_qcm (
                    id_session,
                    id_qcm,
                    id_etudiant,
                    id_candidat,
                    tentative_no,
                    numero_etudiant_detecte,
                    matricule_candidat_detecte,
                    nom_ocr,
                    image_path,
                    overlay_path,
                    raw_answers,
                    raw_id_detection,
                    score,
                    max_score,
                    methode_reponses,
                    methode_id,
                    modele_reponses_version,
                    modele_id_version,
                    confidence_reponses,
                    confidence_id,
                    statut_scan
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s::jsonb, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id_instance
                """,
                (
                    session_ctx["id_session"],
                    session_ctx["id_qcm"],
                    participant.get("id_etudiant"),
                    participant.get("id_candidat"),
                    tentative_no,
                    participant.get("detected_student"),
                    participant.get("detected_candidate"),
                    participant.get("detected_name"),
                    payload.get("image_path"),
                    payload.get("overlay_path"),
                    json.dumps(answers, ensure_ascii=False),
                    json.dumps(raw_id_detection, ensure_ascii=False),
                    score,
                    max_score,
                    method_used,
                    str(payload.get("id_method") or payload.get("methode_id") or "cnn_id"),
                    str(model_versions.get("answers") or payload.get("answers_model_version") or ""),
                    str(model_versions.get("id") or payload.get("id_model_version") or ""),
                    payload.get("confidence_answers"),
                    payload.get("confidence_id"),
                    status,
                ),
            )
            instance_id = int(cur.fetchone()["id_instance"])
            for question in questions:
                numero = int(question["numero"])
                cur.execute(
                    """
                    INSERT INTO reponse_soumise (
                        id_instance,
                        id_question,
                        est_correcte,
                        score_obtenu,
                        confidence
                    )
                    VALUES (%s, %s, %s, %s, %s)
                    RETURNING id_reponse_soumise
                    """,
                    (
                        instance_id,
                        question["id_question"],
                        numero in correct_questions,
                        score_by_num.get(numero, 0.0),
                        None,
                    ),
                )
                submission_answer_id = int(cur.fetchone()["id_reponse_soumise"])
                for label in answers.get(str(numero), []):
                    choice = question["choices"].get(label)
                    if not choice:
                        continue
                    cur.execute(
                        """
                        INSERT INTO reponse_soumise_choix (id_reponse_soumise, id_reponse)
                        VALUES (%s, %s)
                        ON CONFLICT DO NOTHING
                        """,
                        (submission_answer_id, choice["id_reponse"]),
                    )
        conn.commit()
    return instance_id


def _as_iso(value) -> str | None:
    return value.isoformat() if hasattr(value, "isoformat") else None


def _fetch_session_results(session_identifier: str) -> list[dict]:
    session_id = _extract_numeric_id(session_identifier)
    query = """
        SELECT
            i.id_instance,
            i.id_session,
            i.id_qcm,
            i.tentative_no,
            i.score,
            i.max_score,
            i.statut_scan::text AS statut_scan,
            i.methode_reponses,
            i.methode_id,
            i.modele_reponses_version,
            i.modele_id_version,
            i.raw_answers,
            i.raw_id_detection,
            i.numero_etudiant_detecte,
            i.matricule_candidat_detecte,
            i.nom_ocr,
            i.prenom_ocr,
            i.created_at,
            i.updated_at,
            e.id_etudiant,
            e.numero_etudiant,
            e.nom AS etudiant_nom,
            e.prenom AS etudiant_prenom,
            e.email AS etudiant_email,
            ca.id_candidat,
            ca.matricule AS candidat_matricule,
            ca.nom AS candidat_nom,
            ca.prenom AS candidat_prenom,
            ca.email AS candidat_email,
            q_totals.total_questions,
            q_totals.max_score AS qcm_max_score,
            COUNT(rs.id_reponse_soumise) FILTER (WHERE rs.est_correcte) AS correct_questions
        FROM instance_qcm i
        JOIN (
            SELECT id_qcm, COUNT(*) AS total_questions, COALESCE(SUM(ponderation), 0) AS max_score
            FROM question
            GROUP BY id_qcm
        ) q_totals ON q_totals.id_qcm = i.id_qcm
        LEFT JOIN etudiant e ON e.id_etudiant = i.id_etudiant
        LEFT JOIN candidat ca ON ca.id_candidat = i.id_candidat
        LEFT JOIN reponse_soumise rs ON rs.id_instance = i.id_instance
        WHERE i.id_session = %s
        GROUP BY
            i.id_instance,
            q_totals.total_questions,
            q_totals.max_score,
            e.id_etudiant,
            ca.id_candidat
        ORDER BY i.created_at DESC, i.id_instance DESC
    """
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(query, (session_id,))
            rows = cur.fetchall()
            instance_ids = [int(row["id_instance"]) for row in rows]
            details_by_instance: dict[int, list[dict]] = {instance_id: [] for instance_id in instance_ids}
            if instance_ids:
                cur.execute(
                    """
                    SELECT
                        rs.id_instance,
                        q.numero,
                        q.ponderation,
                        rs.est_correcte,
                        rs.score_obtenu,
                        COALESCE(
                            array_agg(DISTINCT selected.label) FILTER (WHERE selected.label IS NOT NULL),
                            ARRAY[]::text[]
                        ) AS detected_answers,
                        COALESCE(
                            array_agg(DISTINCT expected.label) FILTER (WHERE expected.label IS NOT NULL),
                            ARRAY[]::text[]
                        ) AS expected_answers
                    FROM reponse_soumise rs
                    JOIN question q ON q.id_question = rs.id_question
                    LEFT JOIN reponse_soumise_choix rsc ON rsc.id_reponse_soumise = rs.id_reponse_soumise
                    LEFT JOIN reponse selected ON selected.id_reponse = rsc.id_reponse
                    LEFT JOIN reponse expected ON expected.id_question = q.id_question AND expected.est_correcte = TRUE
                    WHERE rs.id_instance = ANY(%s)
                    GROUP BY rs.id_instance, q.numero, q.ponderation, rs.est_correcte, rs.score_obtenu
                    ORDER BY rs.id_instance, q.numero
                    """,
                    (instance_ids,),
                )
                for detail in cur.fetchall():
                    detected = sorted(list(detail.get("detected_answers") or []))
                    expected = sorted(list(detail.get("expected_answers") or []))
                    details_by_instance.setdefault(int(detail["id_instance"]), []).append(
                        {
                            "question": int(detail["numero"]),
                            "detected": detected,
                            "expected": expected,
                            "is_correct": bool(detail["est_correcte"]),
                            "score": float(detail["score_obtenu"] or 0),
                            "points": float(detail["ponderation"] or 0),
                        }
                    )

    results: list[dict] = []
    for row in rows:
        student = None
        if row.get("id_etudiant") is not None:
            student = {
                "type": "student",
                "id": int(row["id_etudiant"]),
                "number": row.get("numero_etudiant"),
                "first_name": row.get("etudiant_prenom"),
                "last_name": row.get("etudiant_nom"),
                "email": row.get("etudiant_email"),
            }
        elif row.get("id_candidat") is not None:
            student = {
                "type": "candidate",
                "id": int(row["id_candidat"]),
                "number": row.get("candidat_matricule"),
                "first_name": row.get("candidat_prenom"),
                "last_name": row.get("candidat_nom"),
                "email": row.get("candidat_email"),
            }

        results.append(
            {
                "submission_id": f"submission-{int(row['id_instance'])}",
                "id_instance": int(row["id_instance"]),
                "session_id": f"session-{int(row['id_session'])}",
                "qcm_id": f"qcm-{int(row['id_qcm'])}",
                "attempt_no": int(row["tentative_no"]),
                "score": float(row["score"] or 0),
                "max_score": float(row["qcm_max_score"] or row["max_score"] or 0),
                "status": row.get("statut_scan"),
                "method_answers": row.get("methode_reponses"),
                "method_id": row.get("methode_id"),
                "answers_model_version": row.get("modele_reponses_version"),
                "id_model_version": row.get("modele_id_version"),
                "answers": row.get("raw_answers") or {},
                "raw_id_detection": row.get("raw_id_detection") or {},
                "detected": {
                    "student_number": row.get("numero_etudiant_detecte"),
                    "candidate_number": row.get("matricule_candidat_detecte"),
                    "name": " ".join(
                        str(part).strip()
                        for part in (row.get("prenom_ocr"), row.get("nom_ocr"))
                        if part
                    ).strip() or row.get("nom_ocr"),
                },
                "student": student,
                "correct_questions": int(row["correct_questions"] or 0),
                "total_questions": int(row["total_questions"] or 0),
                "answer_details": details_by_instance.get(int(row["id_instance"]), []),
                "created_at": _as_iso(row.get("created_at")),
                "updated_at": _as_iso(row.get("updated_at")),
            }
        )
    return results


def _format_score(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:.2f}".rstrip("0").rstrip(".")


def _score_out_of_twenty(score: float, max_score: float) -> float:
    if max_score <= 0:
        return 0.0
    return round((score / max_score) * 20, 2)


def _result_email_payload(session_ctx: dict, qcm_detail: dict, result: dict) -> dict:
    student = result.get("student") or {}
    first_name = student.get("first_name") or ""
    last_name = student.get("last_name") or ""
    full_name = " ".join(str(part).strip() for part in (first_name, last_name) if part).strip() or "etudiant"
    score = float(result.get("score") or 0)
    max_score = float(result.get("max_score") or 0)
    score_20 = _score_out_of_twenty(score, max_score)
    subject_name = qcm_detail.get("subject", {}).get("name") or "la matiere"
    exam_title = qcm_detail.get("title") or session_ctx.get("qcm_title") or session_ctx.get("libelle")
    class_name = session_ctx.get("classe_nom") or qcm_detail.get("class", {}).get("name") or ""
    raw_score = f"{_format_score(score)}/{_format_score(max_score)}"
    score_20_text = f"{_format_score(score_20)}/20"
    email_subject = f"Resultat QCM - {subject_name} - {exam_title}"
    preheader = f"Votre note: {score_20_text} ({raw_score})."
    html = f"""
    <!doctype html>
    <html>
      <body style="margin:0;background:#f3f6fb;font-family:Segoe UI,Arial,sans-serif;color:#243044;">
        <table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="background:#f3f6fb;padding:28px 0;">
          <tr>
            <td align="center">
              <table role="presentation" width="640" cellspacing="0" cellpadding="0" style="max-width:640px;width:100%;background:#ffffff;border:1px solid #d9e0ea;border-radius:10px;overflow:hidden;">
                <tr>
                  <td style="padding:22px 26px;border-bottom:1px solid #d9e0ea;">
                    <img src="cid:junia-logo" alt="JUNIA Maroc" style="max-width:190px;height:auto;display:block;margin-bottom:14px;" />
                    <div style="font-size:12px;letter-spacing:3px;text-transform:uppercase;color:#315987;font-weight:800;">QCM Corrector</div>
                    <h1 style="margin:8px 0 0;color:#111827;font-size:24px;">Resultat de correction</h1>
                  </td>
                </tr>
                <tr>
                  <td style="padding:26px;">
                    <p style="margin:0 0 14px;font-size:16px;">Bonjour {full_name},</p>
                    <p style="margin:0 0 18px;line-height:1.55;">Votre copie a ete corrigee pour <strong>{exam_title}</strong>.</p>
                    <table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="border-collapse:collapse;margin:18px 0;border:1px solid #d9e0ea;">
                      <tr><td style="padding:12px;border-bottom:1px solid #d9e0ea;color:#657386;">Matiere</td><td style="padding:12px;border-bottom:1px solid #d9e0ea;font-weight:700;">{subject_name}</td></tr>
                      <tr><td style="padding:12px;border-bottom:1px solid #d9e0ea;color:#657386;">Controle / concours</td><td style="padding:12px;border-bottom:1px solid #d9e0ea;font-weight:700;">{exam_title}</td></tr>
                      <tr><td style="padding:12px;border-bottom:1px solid #d9e0ea;color:#657386;">Classe</td><td style="padding:12px;border-bottom:1px solid #d9e0ea;font-weight:700;">{class_name or "-"}</td></tr>
                      <tr><td style="padding:12px;color:#657386;">Note</td><td style="padding:12px;font-weight:900;color:#203a63;font-size:22px;">{score_20_text} <span style="font-size:14px;color:#657386;">({raw_score})</span></td></tr>
                    </table>
                    <p style="margin:18px 0 0;color:#657386;line-height:1.5;">Ce message est genere automatiquement depuis la plateforme QCM Corrector.</p>
                  </td>
                </tr>
              </table>
              <div style="display:none;max-height:0;overflow:hidden;">{preheader}</div>
            </td>
          </tr>
        </table>
      </body>
    </html>
    """
    text = (
        f"Bonjour {full_name},\n\n"
        f"Votre copie a ete corrigee pour {exam_title}.\n"
        f"Matiere: {subject_name}\n"
        f"Classe: {class_name or '-'}\n"
        f"Note: {score_20_text} ({raw_score})\n\n"
        "Ce message est genere automatiquement depuis QCM Corrector."
    )
    return {
        "to": student.get("email"),
        "student_name": full_name,
        "student_number": student.get("number"),
        "subject": email_subject,
        "html": html,
        "text": text,
        "score": score,
        "max_score": max_score,
        "score_out_of_20": score_20,
        "raw_score_label": raw_score,
        "score_20_label": score_20_text,
    }


def _send_result_email(message_payload: dict) -> None:
    smtp = _smtp_settings()
    message = MIMEMultipart("related")
    message["Subject"] = message_payload["subject"]
    message["From"] = smtp["from"]
    message["To"] = message_payload["to"]

    alternative = MIMEMultipart("alternative")
    alternative.attach(MIMEText(message_payload["text"], "plain", "utf-8"))
    alternative.attach(MIMEText(message_payload["html"], "html", "utf-8"))
    message.attach(alternative)

    if LOGO_PATH.exists():
        logo = MIMEImage(LOGO_PATH.read_bytes())
        logo.add_header("Content-ID", "<junia-logo>")
        logo.add_header("Content-Disposition", "inline", filename="logo-junia.png")
        message.attach(logo)

    if smtp["use_tls"]:
        context = ssl.create_default_context()
        with smtplib.SMTP(smtp["host"], smtp["port"], timeout=20) as server:
            server.starttls(context=context)
            if smtp["user"]:
                server.login(smtp["user"], smtp["password"])
            server.send_message(message)
    else:
        with smtplib.SMTP(smtp["host"], smtp["port"], timeout=20) as server:
            if smtp["user"]:
                server.login(smtp["user"], smtp["password"])
            server.send_message(message)


def _set_session_email_status(session_id: int, status: str) -> None:
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE session_correction SET email_status = %s, updated_at = now() WHERE id_session = %s",
                (status, session_id),
            )
        conn.commit()


def _cnn_all_marked(answers: dict[str, list[str]], choices_count: int, min_ratio: float = 0.85) -> bool:
    if not answers or choices_count <= 0:
        return False
    all_marked = 0
    total = 0
    for vals in answers.values():
        if vals is None:
            continue
        total += 1
        if isinstance(vals, (list, tuple, set)) and len(vals) >= choices_count:
            all_marked += 1
    if total == 0:
        return False
    return (all_marked / total) >= min_ratio


def normalize_answer_key(raw: dict | None, choices: list[str], questions: int) -> dict[str, list[str]]:
    if not raw:
        return {}
    choice_set = set(choices)
    norm: dict[str, list[str]] = {}
    for i in range(1, questions + 1):
        key = str(i)
        val = raw.get(key) if isinstance(raw, dict) else None
        if val is None:
            continue
        collected: list[str] = []
        values = val if isinstance(val, (list, tuple, set)) else [val]
        for item in values:
            if item is None:
                continue
            text = str(item).strip().upper()
            for ch in text:
                if ch in choice_set and ch not in collected:
                    collected.append(ch)
        if collected:
            # keep choices order
            ordered = [c for c in choices if c in collected]
            norm[key] = ordered
    return norm


def normalize_statements(raw: dict | None, questions: int) -> dict[str, str]:
    if not raw:
        return {}
    norm: dict[str, str] = {}
    for i in range(1, questions + 1):
        key = str(i)
        val = raw.get(key) if isinstance(raw, dict) else None
        if val is None:
            continue
        text = str(val).strip()
        if text:
            norm[key] = text
    return norm


def normalize_choices_text(raw: dict | None, questions: int) -> dict[str, dict[str, str]]:
    if not raw:
        return {}
    norm: dict[str, dict[str, str]] = {}
    for i in range(1, questions + 1):
        key = str(i)
        row = raw.get(key) if isinstance(raw, dict) else None
        if not isinstance(row, dict):
            continue
        row_norm: dict[str, str] = {}
        for letter in ("A", "B", "C", "D"):
            val = row.get(letter)
            if val is None:
                continue
            text = str(val).strip()
            if text:
                row_norm[letter] = text
        if row_norm:
            norm[key] = row_norm
    return norm


def load_answer_data(cfg: dict) -> tuple[dict[str, list[str]], dict[str, str], dict[str, dict[str, str]], str]:
    if ANSWER_KEY_PATH.exists():
        with open(ANSWER_KEY_PATH, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
        if isinstance(data, dict) and ("answers" in data or "statements" in data or "choices_text" in data):
            answers = normalize_answer_key(data.get("answers", {}), cfg["choices"], cfg["questions"])
            statements = normalize_statements(data.get("statements", {}), cfg["questions"])
            choices_text = normalize_choices_text(data.get("choices_text", {}), cfg["questions"])
            return answers, statements, choices_text, "file"
        answers = normalize_answer_key(data if isinstance(data, dict) else {}, cfg["choices"], cfg["questions"])
        return answers, {}, {}, "file"

    answers = normalize_answer_key(cfg.get("answer_key", {}), cfg["choices"], cfg["questions"])
    statements = normalize_statements(cfg.get("statements", {}), cfg["questions"]) if isinstance(cfg, dict) else {}
    choices_text = normalize_choices_text(cfg.get("choices_text", {}), cfg["questions"]) if isinstance(cfg, dict) else {}
    return answers, statements, choices_text, "template"


@router.get("/health")
def health() -> dict:
    return {"status": "ok"}


@router.get("/api/health/db")
def health_db() -> dict:
    try:
        return ping_database()
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"{type(exc).__name__}: {exc}") from exc


@router.get("/template")
def template() -> dict:
    return load_template()


@router.get("/answer-key")
def get_answer_key() -> dict:
    cfg = load_template()
    answers, statements, choices_text, source = load_answer_data(cfg)
    return {"answers": answers, "statements": statements, "choices_text": choices_text, "source": source}


@router.put("/answer-key")
def set_answer_key(payload: dict = Body(...)) -> dict:
    cfg = load_template()
    raw_answers = payload.get("answers", payload)
    raw_statements = payload.get("statements", {})
    raw_choices = payload.get("choices_text", {})
    answers = normalize_answer_key(raw_answers, cfg["choices"], cfg["questions"])
    statements = normalize_statements(raw_statements, cfg["questions"])
    choices_text = normalize_choices_text(raw_choices, cfg["questions"])
    with open(ANSWER_KEY_PATH, "w", encoding="utf-8") as f:
        json.dump({"answers": answers, "statements": statements, "choices_text": choices_text}, f, ensure_ascii=False, indent=2)
    return {
        "status": "ok",
        "saved": len(answers),
        "statements": len(statements),
        "choices_text": len(choices_text),
    }


@router.post("/prof/template")
async def upload_template(template: UploadFile = File(...)) -> dict:
    filename = (template.filename or "").lower()
    ext = Path(filename).suffix
    if ext not in (".pdf", ".jpg", ".jpeg", ".png"):
        raise HTTPException(status_code=400, detail="Format invalide (pdf/jpg/png).")
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    if ext == ".pdf":
        if TEMPLATE_PDF_PATH.exists():
            TEMPLATE_PDF_PATH.unlink()
        target = TEMPLATE_PDF_PATH
    else:
        # normalize jpeg -> jpg
        if ext == ".jpeg":
            ext = ".jpg"
        for existing in TEMPLATE_DIR.glob("template.jpg"):
            existing.unlink()
        for existing in TEMPLATE_DIR.glob("template.png"):
            existing.unlink()
        target = TEMPLATE_DIR / f"template{ext}"
    content = await template.read()
    with open(target, "wb") as f:
        f.write(content)
    return {"status": "ok", "file": target.name}

@router.get("/prof/template.pdf")
def download_template_pdf() -> FileResponse:
    template_pdf = get_template_pdf_path()
    if not template_pdf:
        raise HTTPException(status_code=404, detail="Template PDF introuvable")
    return FileResponse(template_pdf, media_type="application/pdf", filename="questionnaire_etudiant.pdf")


@router.post("/prof/questionnaire.pdf")
def generate_pdf(payload: dict = Body(default={})):  # noqa: B006
    cfg = load_template()
    raw_answers = payload.get("answers") if isinstance(payload, dict) else None
    raw_statements = payload.get("statements") if isinstance(payload, dict) else None
    raw_choices = payload.get("choices_text") if isinstance(payload, dict) else None
    if raw_answers or raw_statements or raw_choices:
        answers = normalize_answer_key(raw_answers or {}, cfg["choices"], cfg["questions"])
        statements = normalize_statements(raw_statements or {}, cfg["questions"])
        choices_text = normalize_choices_text(raw_choices or {}, cfg["questions"])
        with open(ANSWER_KEY_PATH, "w", encoding="utf-8") as f:
            json.dump({"answers": answers, "statements": statements, "choices_text": choices_text}, f, ensure_ascii=False, indent=2)
    else:
        answers, statements, choices_text, _ = load_answer_data(cfg)

    title = payload.get("title", "EXAMEN QCM") if isinstance(payload, dict) else "EXAMEN QCM"
    mode = payload.get("mode", "blank") if isinstance(payload, dict) else "blank"
    style = payload.get("style", "sheet") if isinstance(payload, dict) else "sheet"
    show_answers = mode == "corrected"
    include_statements = bool(payload.get("include_statements", False)) if isinstance(payload, dict) else False
    if mode == "blank" and not include_statements:
        template_pdf = get_template_pdf_path()
        if template_pdf:
            return FileResponse(template_pdf, media_type="application/pdf", filename="questionnaire_etudiant.pdf")
    use_template = bool(payload.get("use_template", False)) if isinstance(payload, dict) else False
    tpl_type, tpl_path = get_template_path() if use_template else (None, None)
    pdf_bytes = generate_questionnaire_pdf(
        answers,
        title=title,
        show_answers=show_answers,
        style=style,
        statements=statements if include_statements else None,
        choices_text=choices_text if include_statements else None,
        template_pdf_path=tpl_path if tpl_type == "pdf" else None,
        template_image_path=tpl_path if tpl_type == "image" else None,
        grid_config=cfg.get("grid"),
    )
    return Response(content=pdf_bytes, media_type="application/pdf")


@router.post("/prof/statements.pdf")
def generate_statements(payload: dict = Body(default={})):  # noqa: B006
    cfg = load_template()
    raw_statements = payload.get("statements") if isinstance(payload, dict) else None
    raw_choices = payload.get("choices_text") if isinstance(payload, dict) else None
    if raw_statements or raw_choices:
        statements = normalize_statements(raw_statements or {}, cfg["questions"])
        choices_text = normalize_choices_text(raw_choices or {}, cfg["questions"])
    else:
        _answers, statements, choices_text, _ = load_answer_data(cfg)
    pdf_bytes = generate_statements_pdf(statements, choices_text)
    return Response(content=pdf_bytes, media_type="application/pdf")


@router.get("/debug/overlay")
def debug_overlay() -> FileResponse:
    sample_path = _get_last_upload_path()
    if not sample_path:
        raise HTTPException(status_code=404, detail="Aucun fichier last_upload.*")
    cfg = load_template()
    data = sample_path.read_bytes()
    warped = preprocess_sheet(data, cfg["paper"]["width"], cfg["paper"]["height"])
    canvas = cv2.cvtColor(warped, cv2.COLOR_GRAY2BGR)
    boxes = compute_boxes(warped, cfg["grid"], cfg["questions"], cfg["choices"])
    for options in boxes.values():
        for (x0, y0, x1, y1) in options.values():
            cv2.rectangle(canvas, (x0, y0), (x1, y1), (0, 255, 255), 1)
    DEBUG_DIR.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(DEBUG_OVERLAY_PATH), canvas)
    return FileResponse(DEBUG_OVERLAY_PATH, media_type="image/jpeg", filename="debug_overlay_last.jpg")

@router.get("/debug/ratios")
def debug_ratios() -> dict:
    sample_path = _get_last_upload_path()
    if not sample_path:
        raise HTTPException(status_code=404, detail="Aucun fichier last_upload.*")
    cfg = load_template()
    data = sample_path.read_bytes()
    warped = preprocess_sheet(data, cfg["paper"]["width"], cfg["paper"]["height"])
    rois = extract_rois(warped, cfg["grid"], cfg["questions"], cfg["choices"])
    inner_margin = float(cfg["marking"].get("inner_margin", 0.12))
    ratios_by_q: dict[int, dict[str, float]] = {}
    for q, options in rois.items():
        ratios_by_q[q] = {choice: black_ratio(patch, inner_margin=inner_margin) for choice, patch in options.items()}
    stats: dict[str, dict[str, float]] = {}
    for choice in cfg["choices"]:
        vals = [ratios.get(choice, 0.0) for ratios in ratios_by_q.values()]
        if not vals:
            stats[choice] = {"min": 0.0, "max": 0.0, "mean": 0.0}
        else:
            stats[choice] = {
                "min": float(min(vals)),
                "max": float(max(vals)),
                "mean": float(sum(vals) / len(vals)),
            }
    return {
        "inner_margin": inner_margin,
        "stats": stats,
    }


@router.get("/api/sessions/{session_id}/results")
def get_session_results(
    session_id: str,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    session_ctx = _fetch_session_context(session_id)
    _assert_teacher_can_access_session(viewer, session_ctx)
    results = _fetch_session_results(session_id)
    return {
        "session": {
            "id": f"session-{int(session_ctx['id_session'])}",
            "label": session_ctx["libelle"],
            "status": session_ctx["statut"],
            "email_status": session_ctx["email_status"],
            "qcm_id": f"qcm-{int(session_ctx['id_qcm'])}",
            "qcm_code": session_ctx["qcm_code"],
            "qcm_title": session_ctx["qcm_title"],
            "class_name": session_ctx.get("classe_nom"),
            "contest_title": session_ctx.get("concours_intitule"),
        },
        "results": results,
        "summary": {
            "copies": len(results),
            "average_score": round(sum(item["score"] for item in results) / len(results), 2) if results else 0,
            "manual_checks": sum(1 for item in results if item["status"] == "a_verifier"),
        },
    }


@router.post("/api/sessions/{session_id}/send-results-email")
def send_session_results_email(
    session_id: str,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
    payload: dict | None = Body(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    session_ctx = _fetch_session_context(session_id)
    _assert_teacher_can_access_session(viewer, session_ctx)
    qcm_detail = _fetch_qcm_detail(int(session_ctx["id_qcm"]))
    reconciled = _reconcile_pending_student_associations(int(session_ctx["id_session"]))
    results = _fetch_session_results(session_id)
    smtp = _smtp_settings()
    smtp_configured = bool(smtp["host"])
    dry_run = bool((payload or {}).get("dry_run", False)) or not smtp_configured

    latest_by_student: dict[str, dict] = {}
    skipped: list[dict] = []
    for result in results:
        student = result.get("student")
        if not student:
            skipped.append({"submission_id": result["submission_id"], "reason": "association_etudiant_manquante"})
            continue
        email_value = str(student.get("email") or "").strip()
        number = str(student.get("number") or result["submission_id"])
        if not email_value:
            skipped.append({"submission_id": result["submission_id"], "reason": "email_etudiant_manquant"})
            continue
        if number not in latest_by_student:
            latest_by_student[number] = result

    prepared = [_result_email_payload(session_ctx, qcm_detail, result) for result in latest_by_student.values()]
    sent: list[dict] = []
    errors: list[dict] = []
    previews: list[dict] = []
    for item in prepared:
        preview = {
            "to": item["to"],
            "student": item["student_name"],
            "student_number": item["student_number"],
            "subject": item["subject"],
            "score": item["score_20_label"],
            "raw_score": item["raw_score_label"],
        }
        if dry_run:
            previews.append(preview)
            continue
        try:
            _send_result_email(item)
            sent.append(preview)
        except Exception as exc:
            errors.append({**preview, "error": f"{type(exc).__name__}: {exc}"})

    if not dry_run:
        if sent and not errors and not skipped:
            _set_session_email_status(int(session_ctx["id_session"]), "envoye")
        elif sent or errors:
            _set_session_email_status(int(session_ctx["id_session"]), "partiel")

    return {
        "status": "preview" if dry_run else "ok" if not errors else "partial",
        "dry_run": dry_run,
        "smtp_configured": smtp_configured,
        "session_id": f"session-{int(session_ctx['id_session'])}",
        "prepared": len(prepared),
        "sent": len(sent),
        "failed": len(errors),
        "reconciled": reconciled,
        "skipped": skipped,
        "previews": previews,
        "sent_items": sent,
        "errors": errors,
        "note": "Configurez SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD et SMTP_FROM dans backend/.env pour envoyer reellement les mails." if dry_run else None,
    }


@router.get("/api/sessions")
def get_sessions_overview(
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    return {"sessions": _fetch_sessions_overview(viewer)}


@router.get("/api/qcms/{qcm_id}")
def get_qcm_detail(
    qcm_id: str,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    qcm_detail = _fetch_qcm_detail(qcm_id)
    _assert_teacher_can_access_qcm(viewer, qcm_detail)
    return qcm_detail


@router.get("/api/qcms/{qcm_id}/subject.pdf")
def download_qcm_subject_pdf(
    qcm_id: str,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> Response:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    qcm_detail = _fetch_qcm_detail(qcm_id)
    _assert_teacher_can_access_qcm(viewer, qcm_detail)
    pdf_bytes = _render_web_page_pdf(
        f"/qcms/{quote(str(qcm_id), safe='')}/subject?pdf=1",
        f"{qcm_detail.get('code') or qcm_id}-sujet.pdf",
    )
    filename = f"{qcm_detail.get('code') or qcm_id}-sujet.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{filename}"'},
    )


@router.get("/api/qcms/{qcm_id}/answer-key.pdf")
def download_qcm_answer_key_pdf(
    qcm_id: str,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> Response:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    qcm_detail = _fetch_qcm_detail(qcm_id)
    _assert_teacher_can_access_qcm(viewer, qcm_detail)
    pdf_bytes = _render_web_page_pdf(
        f"/qcms/{quote(str(qcm_id), safe='')}/answer-key?pdf=1",
        f"{qcm_detail.get('code') or qcm_id}-grille.pdf",
    )
    filename = f"{qcm_detail.get('code') or qcm_id}-grille.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{filename}"'},
    )


@router.post("/api/admin/classes")
def create_admin_class(
    payload: dict = Body(...),
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    name = _require_text(payload.get("name") or payload.get("nom"), "nom")
    level = str(payload.get("level") or payload.get("niveau") or name).strip() or name
    academic_year = _normalize_academic_year(payload.get("academic_year") or payload.get("annee_scolaire"))
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO classe (nom, niveau, annee_scolaire)
                    VALUES (%s, %s, %s)
                    RETURNING id_classe
                    """,
                    (name, level, academic_year),
                )
                class_id = int(cur.fetchone()["id_classe"])
            conn.commit()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Creation classe impossible: {exc}") from exc
    return {"status": "ok", "class_id": class_id, "classes": _fetch_classes_with_students(viewer)}


@router.patch("/api/admin/classes/{class_id}")
def update_admin_class(
    class_id: int,
    payload: dict = Body(...),
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    name = _require_text(payload.get("name") or payload.get("nom"), "nom")
    level = str(payload.get("level") or payload.get("niveau") or name).strip() or name
    academic_year = _normalize_academic_year(payload.get("academic_year") or payload.get("annee_scolaire"))
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    UPDATE classe
                    SET nom = %s, niveau = %s, annee_scolaire = %s, updated_at = now()
                    WHERE id_classe = %s
                    RETURNING id_classe
                    """,
                    (name, level, academic_year, class_id),
                )
                if not cur.fetchone():
                    raise HTTPException(status_code=404, detail="Classe introuvable.")
            conn.commit()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Modification classe impossible: {exc}") from exc
    return {"status": "ok", "classes": _fetch_classes_with_students(viewer)}


@router.delete("/api/admin/classes/{class_id}")
def delete_admin_class(
    class_id: int,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute("DELETE FROM classe WHERE id_classe = %s RETURNING id_classe", (class_id,))
                if not cur.fetchone():
                    raise HTTPException(status_code=404, detail="Classe introuvable.")
            conn.commit()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=409, detail="Suppression impossible: cette classe est deja liee a des etudiants, matieres, QCM ou sessions. Supprimez/deplacez les elements lies avant de supprimer la classe.") from exc
    return {"status": "ok", "classes": _fetch_classes_with_students(viewer)}


@router.post("/api/admin/classes/{class_id}/students")
def create_admin_student(
    class_id: int,
    payload: dict = Body(...),
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                student_id = _insert_student(cur, payload, class_id)
                linked_submissions = _link_pending_instances_for_student(cur, student_id)
            conn.commit()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Ajout etudiant impossible: {exc}") from exc
    return {
        "status": "ok",
        "student_id": student_id,
        "linked_submissions": linked_submissions,
        "classes": _fetch_classes_with_students(viewer),
    }


@router.patch("/api/admin/students/{student_id}")
def update_admin_student(
    student_id: int,
    payload: dict = Body(...),
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    class_id = int(payload.get("class_id") or payload.get("id_classe") or 0)
    student = _normalize_student_payload(payload, class_id)
    active = bool(payload.get("active", payload.get("actif", True)))
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT id_personne FROM etudiant WHERE id_etudiant = %s", (student_id,))
                row = cur.fetchone()
                if not row:
                    raise HTTPException(status_code=404, detail="Etudiant introuvable.")
                person_id = row.get("id_personne")
                cur.execute(
                    """
                    UPDATE etudiant
                    SET numero_etudiant = %s, nom = %s, prenom = %s, email = %s, id_classe = %s, actif = %s, updated_at = now()
                    WHERE id_etudiant = %s
                    """,
                    (student["number"], student["last_name"], student["first_name"], student["email"], class_id, active, student_id),
                )
                if person_id:
                    cur.execute(
                        """
                        UPDATE personne
                        SET numero_etudiant = %s, nom = %s, prenom = %s, email = %s, id_groupe = %s, updated_at = now()
                        WHERE id_personne = %s
                        """,
                        (student["number"], student["last_name"], student["first_name"], student["email"], class_id, person_id),
                    )
                linked_submissions = _link_pending_instances_for_student(cur, student_id)
            conn.commit()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Modification etudiant impossible: {exc}") from exc
    return {"status": "ok", "linked_submissions": linked_submissions, "classes": _fetch_classes_with_students(viewer)}


@router.delete("/api/admin/students/{student_id}")
def delete_admin_student(
    student_id: int,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    mode = "deleted"
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT id_personne FROM etudiant WHERE id_etudiant = %s", (student_id,))
                row = cur.fetchone()
                if not row:
                    raise HTTPException(status_code=404, detail="Etudiant introuvable.")
                person_id = row.get("id_personne")
                cur.execute("DELETE FROM etudiant WHERE id_etudiant = %s", (student_id,))
                if person_id:
                    cur.execute("DELETE FROM personne WHERE id_personne = %s", (person_id,))
            conn.commit()
    except HTTPException:
        raise
    except Exception:
        mode = "deactivated"
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE etudiant SET actif = FALSE, updated_at = now() WHERE id_etudiant = %s", (student_id,))
                if cur.rowcount == 0:
                    raise HTTPException(status_code=404, detail="Etudiant introuvable.")
            conn.commit()
    return {"status": "ok", "mode": mode, "classes": _fetch_classes_with_students(viewer)}


@router.post("/api/admin/classes/import")
def import_admin_class(
    name: str = Form(...),
    level: str | None = Form(default=None),
    academic_year: str | None = Form(default=None),
    file: UploadFile = File(...),
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    class_name = _require_text(name, "nom")
    class_level = str(level or class_name).strip() or class_name
    year = _normalize_academic_year(academic_year)
    raw = file.file.read()
    rows = _csv_rows_from_upload(raw, file.filename or "import.csv")
    imported = 0
    skipped: list[dict] = []
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO classe (nom, niveau, annee_scolaire)
                    VALUES (%s, %s, %s)
                    ON CONFLICT (nom, annee_scolaire)
                    DO UPDATE SET niveau = EXCLUDED.niveau, updated_at = now()
                    RETURNING id_classe
                    """,
                    (class_name, class_level, year),
                )
                class_id = int(cur.fetchone()["id_classe"])
                for index, row in enumerate(rows, start=2):
                    try:
                        student_id = _insert_student(cur, _student_payload_from_import_row(row), class_id)
                        _link_pending_instances_for_student(cur, student_id)
                        imported += 1
                    except Exception as exc:
                        skipped.append({"line": index, "reason": str(exc)})
            conn.commit()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Import classe impossible: {exc}") from exc
    return {"status": "ok", "imported": imported, "skipped": skipped, "classes": _fetch_classes_with_students(viewer)}


@router.get("/api/admin/classes")
def get_admin_classes(
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    return {"classes": _fetch_classes_with_students(viewer)}


@router.get("/api/admin/profs")
def get_admin_professors(
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    if not viewer or viewer["role"] != "admin":
        raise HTTPException(status_code=403, detail="Gestion des professeurs reservee a l'admin.")
    return {"professors": _fetch_professors()}


@router.post("/api/admin/profs")
def create_admin_professor(
    payload: dict = Body(...),
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    data = _normalize_professor_payload(payload)
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT id_utilisateur, role::text AS role FROM utilisateur WHERE lower(email) = %s LIMIT 1",
                    (data["email"],),
                )
                existing = cur.fetchone()
                if existing and existing["role"] != "enseignant":
                    raise HTTPException(status_code=409, detail="Cet email est deja utilise par un compte non enseignant.")

                if existing:
                    professor_id = int(existing["id_utilisateur"])
                    cur.execute(
                        """
                        UPDATE utilisateur
                        SET nom = %s, prenom = %s, password_hash = %s, actif = TRUE, changement_mdp_requis = FALSE, updated_at = now()
                        WHERE id_utilisateur = %s
                        """,
                        (data["last_name"], data["first_name"], data["password"], professor_id),
                    )
                else:
                    cur.execute(
                        """
                        INSERT INTO utilisateur (nom, prenom, email, password_hash, role, actif, changement_mdp_requis)
                        VALUES (%s, %s, %s, %s, 'enseignant', TRUE, FALSE)
                        RETURNING id_utilisateur
                        """,
                        (data["last_name"], data["first_name"], data["email"], data["password"]),
                    )
                    professor_id = int(cur.fetchone()["id_utilisateur"])

                cur.execute(
                    """
                    INSERT INTO enseignant (id_user, prenom)
                    VALUES (%s, %s)
                    ON CONFLICT (id_user) DO UPDATE SET prenom = EXCLUDED.prenom
                    """,
                    (professor_id, data["first_name"]),
                )
                subject_id = _upsert_professor_subject(cur, professor_id, data)
            conn.commit()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Creation professeur impossible: {exc}") from exc
    return {
        "status": "ok",
        "professor_id": professor_id,
        "subject_id": subject_id,
        "login": {"email": data["email"], "password": data["password"]},
        "professors": _fetch_professors(),
    }


@router.patch("/api/admin/profs/{professor_id}")
def update_admin_professor(
    professor_id: int,
    payload: dict = Body(...),
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    try:
        with connect() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT id_utilisateur, password_hash
                    FROM utilisateur
                    WHERE id_utilisateur = %s AND role = 'enseignant'
                    LIMIT 1
                    """,
                    (professor_id,),
                )
                row = cur.fetchone()
                if not row:
                    raise HTTPException(status_code=404, detail="Professeur introuvable.")

                data = _normalize_professor_payload(payload, existing_password=row["password_hash"])
                cur.execute(
                    """
                    SELECT id_utilisateur
                    FROM utilisateur
                    WHERE lower(email) = %s AND id_utilisateur <> %s
                    LIMIT 1
                    """,
                    (data["email"], professor_id),
                )
                if cur.fetchone():
                    raise HTTPException(status_code=409, detail="Cet email est deja utilise par un autre compte.")

                cur.execute(
                    """
                    UPDATE utilisateur
                    SET nom = %s, prenom = %s, email = %s, password_hash = %s, actif = %s, changement_mdp_requis = FALSE, updated_at = now()
                    WHERE id_utilisateur = %s
                    """,
                    (data["last_name"], data["first_name"], data["email"], data["password"], data["active"], professor_id),
                )
                cur.execute(
                    """
                    INSERT INTO enseignant (id_user, prenom)
                    VALUES (%s, %s)
                    ON CONFLICT (id_user) DO UPDATE SET prenom = EXCLUDED.prenom
                    """,
                    (professor_id, data["first_name"]),
                )
                raw_subject_id = payload.get("subject_id") or payload.get("id_matiere")
                subject_id = int(raw_subject_id) if raw_subject_id else None
                subject_id = _upsert_professor_subject(cur, professor_id, data, subject_id)
            conn.commit()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Modification professeur impossible: {exc}") from exc
    return {
        "status": "ok",
        "professor_id": professor_id,
        "subject_id": subject_id,
        "login": {"email": data["email"], "password": data["password"]},
        "professors": _fetch_professors(),
    }


@router.delete("/api/admin/profs/{professor_id}")
def delete_admin_professor(
    professor_id: int,
    role: str | None = Query(default=None),
    user_id: str | None = Query(default=None),
    email: str | None = Query(default=None),
) -> dict:
    viewer = _resolve_viewer(role=role, user_id=user_id, email=email)
    _require_admin_viewer(viewer)
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE utilisateur
                SET actif = FALSE, updated_at = now()
                WHERE id_utilisateur = %s AND role = 'enseignant'
                RETURNING id_utilisateur
                """,
                (professor_id,),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="Professeur introuvable.")
        conn.commit()
    return {"status": "ok", "mode": "deactivated", "professors": _fetch_professors()}


@router.post("/api/auth/login")
def login_user(payload: dict = Body(...)) -> dict:
    email = str(payload.get("email") or "").strip().lower()
    password = str(payload.get("password") or "")
    expected_role = str(payload.get("role") or "").strip().lower()
    if not email or not password:
        raise HTTPException(status_code=400, detail="Email et mot de passe requis.")
    if expected_role and expected_role not in {"admin", "enseignant"}:
        raise HTTPException(status_code=400, detail="Role utilisateur invalide.")

    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id_utilisateur, email, nom, prenom, password_hash, role::text AS role, actif
                FROM utilisateur
                WHERE lower(email) = %s
                LIMIT 1
                """,
                (email,),
            )
            row = cur.fetchone()
    if not row or not row["actif"] or row["password_hash"] != password:
        raise HTTPException(status_code=401, detail="Identifiants incorrects.")
    if expected_role and row["role"] != expected_role:
        raise HTTPException(status_code=403, detail="Ce compte ne correspond pas a ce profil.")

    first_name = row["prenom"]
    last_name = row["nom"]
    initials = "".join(part[:1].upper() for part in (first_name, last_name) if part)[:2] or "U"
    return {
        "user": {
            "id": int(row["id_utilisateur"]),
            "email": row["email"],
            "first_name": first_name,
            "last_name": last_name,
            "role": row["role"],
            "initials": initials,
        }
    }


@router.post("/api/qcms/sync")
def sync_qcm_from_web(payload: dict = Body(...)) -> dict:
    return _sync_web_qcm(payload)


@router.post("/api/qcms/draft")
def save_qcm_draft_from_web(payload: dict = Body(...)) -> dict:
    return _save_web_qcm_draft(payload)


def _process_detected_result(session_id: str, payload: dict) -> dict:
    if isinstance(payload.get("error"), dict):
        raise HTTPException(status_code=422, detail=payload["error"])

    session_ctx = _fetch_session_context(session_id)
    raw_answers = _extract_answers_from_detection_payload(payload)
    if not raw_answers:
        raise HTTPException(
            status_code=400,
            detail="Aucune reponse detectee. Envoyez answers ou des champs q1, q2, q3...",
        )

    questions = _fetch_qcm_questions(int(session_ctx["id_qcm"]))
    if not questions:
        raise HTTPException(status_code=404, detail=f"Aucune question trouvee pour le QCM {session_ctx['id_qcm']}.")

    known_question_numbers = {str(question["numero"]) for question in questions}
    ignored_detected_questions = sorted(
        int(key)
        for key in raw_answers.keys()
        if str(key).isdigit() and str(key) not in known_question_numbers
    )
    all_detected_answers = _normalize_detected_answer_values(raw_answers)
    answers = _normalize_detected_answers(raw_answers, questions)
    participant = _match_detected_participant(session_ctx, payload)
    score, max_score, correct_questions, wrong_questions, expected = _compute_weighted_score(answers, questions)
    method_used = str(payload.get("answers_method") or payload.get("method") or "cnn")

    submission_id = _store_detected_result(
        session_ctx=session_ctx,
        payload=payload,
        participant=participant,
        answers=all_detected_answers,
        score=score,
        max_score=max_score,
        correct_questions=correct_questions,
        questions=questions,
        method_used=method_used,
    )

    warnings: list[str] = []
    if not participant.get("matched") and not participant.get("match_warning"):
        warnings.append("Identifiant detecte non trouve dans la base pour cette session.")
    if participant.get("match_warning"):
        warnings.append(str(participant["match_warning"]))
    if ignored_detected_questions:
        warnings.append(
            f"{len(ignored_detected_questions)} reponses detectees ont ete ignorees car ces numeros de questions ne sont pas presents dans ce QCM."
        )

    return {
        "submission_id": f"submission-{submission_id}",
        "session_id": f"session-{int(session_ctx['id_session'])}",
        "qcm_id": f"qcm-{int(session_ctx['id_qcm'])}",
        "raw_detection": payload,
        "session_context": {
            "label": session_ctx["libelle"],
            "qcm_code": session_ctx["qcm_code"],
            "qcm_title": session_ctx["qcm_title"],
            "qcm_type": session_ctx["qcm_type"],
            "qcm_status": session_ctx["qcm_status"],
            "class_name": session_ctx.get("classe_nom"),
            "contest_title": session_ctx.get("concours_intitule"),
        },
        "student": participant.get("matched"),
        "detected": {
            "student_number": participant.get("detected_student"),
            "candidate_number": participant.get("detected_candidate"),
            "name": participant.get("detected_name"),
        },
        "status": "valide" if participant.get("matched") else "a_verifier",
        "score": score,
        "max_score": max_score,
        "answers": all_detected_answers,
        "scored_answers": answers,
        "expected": expected,
        "correct_questions": correct_questions,
        "wrong_questions": wrong_questions,
        "warnings": warnings,
    }


@router.post("/api/sessions/{session_id}/detected-result")
def receive_detected_result(session_id: str, payload: dict = Body(...)) -> dict:
    return _process_detected_result(session_id, payload)


@router.post("/api/detected-result")
def receive_detected_result_auto_session(payload: dict = Body(...)) -> dict:
    session_ctx = _resolve_session_from_detection_payload(payload)
    result = _process_detected_result(str(session_ctx["id_session"]), payload)
    result["auto_resolved_session"] = True
    result["resolved_from"] = "student_id_to_class_to_published_qcm_session"
    return result


def _looks_like_ai_result_batch(items: object) -> bool:
    if not isinstance(items, list):
        return False
    for item in items:
        if not isinstance(item, dict):
            continue
        if any(
            key in item
            for key in (
                "answers",
                "detected_answers",
                "student_id",
                "studentId",
                "student_id_detected",
                "numero_etudiant_detecte",
                "candidate_id",
                "matricule_candidat_detecte",
                "session_id",
                "sessionId",
            )
        ):
            return True
    return False


def _extract_ai_result_items(payload: object) -> tuple[list[object], object | None, bool]:
    if isinstance(payload, list):
        return payload, None, True
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="Le payload doit etre un objet JSON ou une liste JSON.")

    inherited_session_id = payload.get("session_id") or payload.get("sessionId")
    for key in ("corrections", "submissions", "copies"):
        if isinstance(payload.get(key), list):
            return payload[key], inherited_session_id, True

    raw_results = payload.get("results")
    if _looks_like_ai_result_batch(raw_results):
        return raw_results, inherited_session_id, True

    return [payload], inherited_session_id, False


def _receive_one_ai_correction(payload: dict) -> dict:
    session_id = payload.get("session_id") or payload.get("sessionId")
    if session_id:
        result = _process_detected_result(str(session_id), payload)
        result["resolved_from"] = "provided_session_id"
        return result

    session_ctx = _resolve_session_from_detection_payload(payload)
    result = _process_detected_result(str(session_ctx["id_session"]), payload)
    result["auto_resolved_session"] = True
    result["resolved_from"] = "student_id_to_class_to_published_qcm_session"
    return result


@router.post("/api/ai/corrections")
def receive_ai_corrections(
    payload: object = Body(...),
    x_api_key: str | None = Header(default=None, alias="X-API-Key"),
) -> dict:
    _require_ai_results_key(x_api_key)
    items, inherited_session_id, is_batch = _extract_ai_result_items(payload)

    results: list[dict] = []
    errors: list[dict] = []
    for index, raw_item in enumerate(items):
        if not isinstance(raw_item, dict):
            errors.append({"index": index, "status": 400, "detail": "Chaque correction doit etre un objet JSON."})
            continue
        normalized = _normalize_ai_result_payload(raw_item, inherited_session_id)
        try:
            results.append(_receive_one_ai_correction(normalized))
        except HTTPException as exc:
            errors.append({"index": index, "status": exc.status_code, "detail": exc.detail})
        except Exception as exc:
            errors.append({"index": index, "status": 500, "detail": f"{type(exc).__name__}: {exc}"})

    response = {
        "status": "partial" if errors and results else "error" if errors else "ok",
        "received": len(items),
        "saved": len(results),
        "failed": len(errors),
        "results": results,
        "errors": errors,
    }
    if not is_batch and errors:
        first = errors[0]
        raise HTTPException(status_code=int(first["status"]), detail=first["detail"])
    return response


@router.post("/api/sessions/{session_id}/ai-corrections")
def receive_session_ai_corrections(
    session_id: str,
    payload: object = Body(...),
    x_api_key: str | None = Header(default=None, alias="X-API-Key"),
) -> dict:
    _require_ai_results_key(x_api_key)
    items, _inherited_session_id, is_batch = _extract_ai_result_items(payload)

    results: list[dict] = []
    errors: list[dict] = []
    for index, raw_item in enumerate(items):
        if not isinstance(raw_item, dict):
            errors.append({"index": index, "status": 400, "detail": "Chaque correction doit etre un objet JSON."})
            continue
        normalized = _normalize_ai_result_payload(raw_item, session_id)
        normalized["session_id"] = session_id
        try:
            result = _process_detected_result(session_id, normalized)
            result["resolved_from"] = "url_session_id"
            results.append(result)
        except HTTPException as exc:
            errors.append({"index": index, "status": exc.status_code, "detail": exc.detail})
        except Exception as exc:
            errors.append({"index": index, "status": 500, "detail": f"{type(exc).__name__}: {exc}"})

    response = {
        "status": "partial" if errors and results else "error" if errors else "ok",
        "received": len(items),
        "saved": len(results),
        "failed": len(errors),
        "results": results,
        "errors": errors,
    }
    if not is_batch and errors:
        first = errors[0]
        raise HTTPException(status_code=int(first["status"]), detail=first["detail"])
    return response


@router.get("/api/ia/config")
def get_ia_config() -> dict:
    return {
        "ia_api_base_url": IA_API_BASE_URL,
        "ia_analyze_path": IA_ANALYZE_PATH,
        "ia_analyze_url": f"{IA_API_BASE_URL}{IA_ANALYZE_PATH}",
        "timeout_seconds": IA_TIMEOUT_SECONDS,
        "note": "Redemarrez le backend apres modification de backend/.env.",
    }


async def _call_external_ia(file: UploadFile) -> dict:
    filename = file.filename or "scan.jpg"
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in {"jpg", "jpeg", "png", "bmp"}:
        raise HTTPException(status_code=400, detail="Format invalide. Formats acceptes: jpg, jpeg, png, bmp.")

    payload = await file.read()
    if not payload:
        raise HTTPException(status_code=400, detail="Fichier vide.")

    try:
        import requests  # type: ignore
    except ModuleNotFoundError as exc:
        raise HTTPException(
            status_code=500,
            detail="Module requests manquant. Installez les dependances backend avec requirements.txt.",
        ) from exc

    analyze_url = f"{IA_API_BASE_URL}{IA_ANALYZE_PATH}"
    try:
        response = requests.post(
            analyze_url,
            files={"file": (filename, payload, file.content_type or "application/octet-stream")},
            timeout=IA_TIMEOUT_SECONDS,
        )
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"API IA inaccessible: {exc}") from exc

    try:
        ia_payload = response.json()
    except ValueError as exc:
        raise HTTPException(status_code=502, detail="L'API IA n'a pas renvoye un JSON valide.") from exc

    if not response.ok:
        detail = ia_payload.get("error") if isinstance(ia_payload, dict) else None
        raise HTTPException(status_code=response.status_code, detail=detail or ia_payload)
    if not isinstance(ia_payload, dict):
        raise HTTPException(status_code=502, detail="La reponse IA doit etre un objet JSON.")

    return ia_payload


@router.post("/api/ia/analyze-only")
async def analyze_with_external_ia_raw(file: UploadFile = File(...)) -> dict:
    ia_payload = await _call_external_ia(file)
    return {
        "ia_analyze_url": f"{IA_API_BASE_URL}{IA_ANALYZE_PATH}",
        "raw_detection": ia_payload,
    }


@router.post("/api/analyze-with-ia")
async def analyze_with_external_ia_auto_session(file: UploadFile = File(...)) -> dict:
    ia_payload = await _call_external_ia(file)
    session_ctx = _resolve_session_from_detection_payload(ia_payload)
    result = _process_detected_result(str(session_ctx["id_session"]), ia_payload)
    result["auto_resolved_session"] = True
    result["resolved_from"] = "ia_response_student_id_to_class_to_published_qcm_session"
    return result


@router.post("/api/sessions/{session_id}/analyze-with-ia")
async def analyze_with_external_ia(
    session_id: str,
    file: UploadFile = File(...),
) -> dict:
    ia_payload = await _call_external_ia(file)
    return _process_detected_result(session_id, ia_payload)


@router.post("/api/sessions/{session_id}/scan")
async def scan_qcm_for_session(
    session_id: str,
    file: UploadFile = File(...),
    qcm_id: str | None = Form(default=None),
    method: str = Form(default="cnn", pattern="^cnn$"),
    save: bool = Form(default=True),
) -> dict:
    payload = await file.read()
    if not payload:
        raise HTTPException(status_code=400, detail="Image vide")

    session_ctx = _fetch_session_context(session_id)
    expected_qcm_id = int(session_ctx["id_qcm"])
    if qcm_id:
        provided_qcm_id = _extract_numeric_id(qcm_id)
        if provided_qcm_id != expected_qcm_id:
            raise HTTPException(
                status_code=400,
                detail=f"Le QCM fourni ({provided_qcm_id}) ne correspond pas a la session ({expected_qcm_id}).",
            )

    if save:
        _save_last_upload(payload, file.filename)
    saved_image = _save_scan_image(payload, file.filename, int(session_ctx["id_session"]))

    cfg = load_template()
    warped = preprocess_sheet(payload, cfg["paper"]["width"], cfg["paper"]["height"])
    rois = extract_rois(warped, cfg["grid"], cfg["questions"], cfg["choices"])

    try:
        answers, cnn_meta = classify_cnn(rois, str(MODEL_PATH), threshold=0.5, return_meta=True)
    except (FileNotFoundError, RuntimeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    questions = _fetch_qcm_questions(expected_qcm_id)
    if not questions:
        raise HTTPException(status_code=404, detail=f"Aucune question trouvee pour le QCM {expected_qcm_id}.")

    score, max_score, correct_questions, wrong_questions, expected = _compute_weighted_score(answers, questions)
    submission_id = _store_scan_result(
        session_ctx=session_ctx,
        image_path=saved_image,
        answers=answers,
        score=score,
        max_score=max_score,
        correct_questions=correct_questions,
        questions=questions,
        method_used="cnn",
    )

    warnings = [
        "Le modele de detection du numero etudiant n'est pas encore raccorde a cette route.",
        "La copie est en statut a_verifier tant que l'identifiant etudiant n'est pas confirme.",
    ]

    return {
        "submission_id": f"submission-{submission_id}",
        "session_id": f"session-{int(session_ctx['id_session'])}",
        "qcm_id": f"qcm-{expected_qcm_id}",
        "student": {
            "id_detected": None,
            "matched_user_id": None,
            "confidence": None,
        },
        "score": score,
        "max_score": max_score,
        "method": method,
        "method_used": "cnn",
        "answers": answers,
        "expected": expected,
        "correct_questions": correct_questions,
        "wrong_questions": wrong_questions,
        "cnn_meta": cnn_meta,
        "warnings": warnings,
        "overlay_url": "/debug/overlay",
        "saved_image_path": str(saved_image),
        "session_context": {
            "label": session_ctx["libelle"],
            "qcm_code": session_ctx["qcm_code"],
            "qcm_title": session_ctx["qcm_title"],
            "classe_nom": session_ctx.get("classe_nom"),
            "concours_intitule": session_ctx.get("concours_intitule"),
        },
    }


@router.post("/scan-qcm")
async def scan_qcm(
    image: UploadFile = File(...),
    method: str = Query(default="cnn", pattern="^cnn$"),
    save_debug: bool = Query(default=True),
) -> dict:
    payload = await image.read()
    if not payload:
        raise HTTPException(status_code=400, detail="Image vide")
    if save_debug:
        _save_last_upload(payload, image.filename)
    cfg = load_template()
    warped = preprocess_sheet(payload, cfg["paper"]["width"], cfg["paper"]["height"])
    rois = extract_rois(warped, cfg["grid"], cfg["questions"], cfg["choices"])
    cnn_meta = None
    try:
        answers, cnn_meta = classify_cnn(rois, str(MODEL_PATH), threshold=0.5, return_meta=True)
    except (FileNotFoundError, RuntimeError) as e:
        raise HTTPException(status_code=400, detail=str(e)) from e

    expected, _statements, _choices_text, source = load_answer_data(cfg)
    score, correct, wrong = score_answers(answers, expected)
    return {
        "score": score,
        "total": cfg["questions"],
        "method": method,
        "method_used": "cnn",
        "answers": answers,
        "expected": expected,
        "expected_source": source,
        "cnn_meta": cnn_meta,
        "fallback_reason": None,
        "correct_questions": correct,
        "wrong_questions": wrong,
    }


